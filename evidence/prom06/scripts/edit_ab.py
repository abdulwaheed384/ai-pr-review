import sys, copy, docx
from docx.oxml.ns import qn
src,dst=sys.argv[1:3]
d=docx.Document(src)
P=d.paragraphs
URL="https://github.com/abdulwaheed384/ai-pr-review"
done=[]
def find(start):
    m=[p for p in d.paragraphs if p.text.strip().startswith(start)]
    assert len(m)==1,(start,len(m)); return m[0]
def settext(p,text):
    rs=p.runs; rs[0].text=text
    for r in rs[1:]: r._r.getparent().remove(r._r)
def setlead(p,lead,rest):
    rs=p.runs; rs[0].text=lead
    if len(rs)<2: raise SystemExit('no second run')
    rs[1].text=rest; rs[1].bold=False
    for r in rs[2:]: r._r.getparent().remove(r._r)
def replace_in(p,old,new):
    assert old in p.text,(old,p.text[:80])
    for r in p.runs:
        if old in r.text: r.text=r.text.replace(old,new); return
    settext(p,p.text.replace(old,new))
def insert_after(p,text,style=None):
    new=copy.deepcopy(p._p); p._p.addnext(new)
    np=docx.text.paragraph.Paragraph(new,p._parent); settext(np,text)
    if style: np.style=d.styles[style]
    return np
def delete(p): p._p.getparent().remove(p._p)

# 1 public URL
p=find("Supporting evidence is in the appendices."); settext(p,f"Supporting evidence is in the appendices, and the complete artefact is publicly available at {URL}.")
p=find("The pipeline is entirely static.")
insert_after(p,f"The complete artefact is publicly available at {URL}. The main branch holds the workflow, the review runner, the prompt and policy files and the baseline Terraform. Each test case has its own branch, and all raw run evidence, scripts and figures are archived unaltered on the prom06-evidence branch (Appendix D).","Body Text")
# AI-use disclosure: include figures
p=find("Generative AI was both the subject"); replace_in(p,"evidence verification, scoring and drafting","evidence verification, scoring, figure preparation and drafting")
# 3 fifteen cases from the start
settext(find("Twelve test cases were specified"),"Fifteen test cases were specified, each with its ground truth recorded before any tool was run against it (Appendix B). Each isolates one Azure misconfiguration or control scenario, and the set includes secure and legitimately public cases to test over-reporting.")
settext(find("Three cases were added on 13 October"),"Three cases are designed to be read against others in the set:")
settext(find("TC13 isolates a least-privilege violation"),"TC13 isolates a least-privilege violation that can only be recognised by comparing a role assignment with a separate requirements file. TC11 tests a similar idea, but its wildcard role is detectable on its own.")
settext(find("The ground truth for TC13–TC15 was committed"),"The fifteen cases are listed in Table 3.2.")
replace_in(find("Each case was built on its own branch"),"byte-identical to the original test-case commits","byte-identical to the commits used in the earlier scanner-informed runs")
replace_in(find("Explanation and remediation quality were scored"),"before TC02–TC12 were executed","before the independent runs were executed")
for s in ("Both scanners produced identical output","Independently of the scanners, the AI reviewer","The second form is a combination"):
    replace_in(find(s)," in TC01–TC15","")
settext(find("Figure 4.4 compares detection rates"),"Figure 4.4 compares detection rates across the fifteen cases. The scanners together detected 67% of the vulnerable cases and the AI 92% (100% if PARTIAL counts).")
p=find("The gap widened when the three later cases")
replace_in(p,"The gap widened when the three later cases were added, because all three targeted conditions without a scanner rule. This illustrates how strongly such comparisons depend on case selection.","The size of the gap depends strongly on case selection. TC09, TC13 and TC15 target conditions for which neither scanner has a rule, and together they account for the whole difference.")
settext(find("Figure 4.4: Detection rate by tool"),"Figure 4.4: Detection rate by tool across the fifteen test cases")
settext(find("Ground truth was fixed before any tool was run, and for TC13"),"Ground truth was fixed before any tool was run on each case.")
replace_in(find("The most important lesson of this project"),"Every AI result for TC01–TC12 was therefore scanner-informed.","Every AI result from that phase was therefore scanner-informed.")
delete(find("On the original twelve-case set"))
# Table 4.2
for t in d.tables:
    if t.rows[0].cells[0].text=='Measure' and 'TC01–TC12' in t.rows[1].cells[0].text:
        t._tbl.remove(t.rows[1]._tr)
        c=t.rows[1].cells[0]; settext(c.paragraphs[0],"Vulnerable cases detected (n = 12)")
        done.append('table 4.2')
# 5 section 5.2 and 5.4
settext(find("The main shortfall is the human dimension of O4"),"The main shortfall is the human dimension of O4: the effect of AI feedback on developers’ understanding was inferred from the quality of the feedback, not observed directly (Section 5.4).")
setlead(find("No developer sessions."),"No developer sessions."," The PROM05 plan included sessions and interviews with developers. Recruiting participants and arranging sessions proved impractical within the project timescale, so I concentrated the available time on a larger controlled experiment with repeated runs. The study therefore assesses the quality of the AI’s feedback rather than its effect on developers, and it did not require ethical approval.")
setlead(find("Azure deployment."),"No Azure deployment."," The test cases were reviewed as pull requests and never merged or applied, so no deliberately insecure resource was created in Azure. The research questions concern review before deployment, so static analysis of the Terraform was sufficient.")
setlead(find("Additions."),"Additions."," I added repeated runs and remediation testing. Both responded to weaknesses that verification exposed in the original design, not to anything the plan anticipated.")
d.save(dst); print('ok',done)
