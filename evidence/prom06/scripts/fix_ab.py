import sys, docx, re
from docx.shared import Cm, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
src,dst=sys.argv[1],sys.argv[2]
d=docx.Document(src)
log=[]
# 1 A4
for s in d.sections:
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
log.append('A4 set')
# 2 fonts and spacing
def arial(style):
    rpr=style.element.get_or_add_rPr(); rf=rpr.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts'); rpr.insert(0,rf)
    for a in list(rf.attrib): del rf.attrib[a]
    for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),'Arial')
    style.font.size=Pt(12)
S={s.name:s for s in d.styles}
for n in ('Normal','Body Text','First Paragraph','Compact','Block Text','Bibliography'):
    if n in S:
        arial(S[n]); S[n].paragraph_format.line_spacing=1.5
log.append('Body styles Arial 12pt, 1.5 spacing')
# tables single spaced
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs: p.paragraph_format.line_spacing=1.0
# TOC entries single
for n in S:
    if n.lower().startswith('toc'): S[n].paragraph_format.line_spacing=1.0
# 3 heading typos
fixes={'Aim':'Aims',
 'Ethical, Social, Professional, Legal and Security Consideration':'Ethical, Social, Professional, Legal and Security Considerations',
 'AI in Software Security Code Review':'AI in Software Security and Code Review',
 'Experimental Design and Test Cases (TC01-15)':'Experimental Design and Test Cases (TC01–TC15)',
 'Comparative Analysis Across TC01-TC15':'Comparative Analysis Across TC01–TC15',
 'Evaluation of Ethical, Legal. Social, Security and Professional Considerations':'Evaluation of Ethical, Legal, Social, Security and Professional Considerations'}
for p in d.paragraphs:
    if p.style.name.startswith('Heading') and p.text.strip() in fixes and p.runs:
        new=fixes[p.text.strip()]; p.runs[0].text=new
        for r in p.runs[1:]: r.text=''
        log.append(f'heading -> {new}')
# 4 empty captions
n=0
for p in list(d.paragraphs):
    if p.style.name in ('Image Caption','Table Caption','Caption') and not p.text.strip() and not p._p.findall('.//'+qn('w:drawing')):
        p._p.getparent().remove(p._p); n+=1
log.append(f'removed {n} empty captions')
# 5 ask Word to refresh fields (TOC) on open
st=d.settings.element; uf=st.find(qn('w:updateFields'))
if uf is None: uf=OxmlElement('w:updateFields'); st.append(uf)
uf.set(qn('w:val'),'true')
d.save(dst); print('\n'.join(log))
