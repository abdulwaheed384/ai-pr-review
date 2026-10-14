"""Apply AB's 14 Oct review requests to his template file, preserving run formatting.
Usage: python3 edit_v3.py in.docx out.docx changelog.json"""
import sys, copy, json, re
import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

src, dst, logf = sys.argv[1:4]
d = docx.Document(src)
LOG = []          # (section, change, before, after)
FAIL = []
XMLSPACE = '{http://www.w3.org/XML/1998/namespace}space'


# ------------------------------------------------------------------ helpers
def tnodes(p_el):
    return [t for t in p_el.iter(qn('w:t'))]

def ptext(p_el):
    return ''.join(t.text or '' for t in tnodes(p_el))

def replace_in(p_el, old, new):
    """Replace the first occurrence of old in the paragraph text, editing w:t nodes in place
    so the formatting of the first affected run is kept."""
    ts = tnodes(p_el); full = ''.join(t.text or '' for t in ts)
    i = full.find(old)
    if i < 0:
        return False
    j = i + len(old); pos = 0; first = True
    for t in ts:
        s = t.text or ''; a, b = pos, pos + len(s); pos = b
        if b <= i or a >= j:
            continue
        lo, hi = max(i, a) - a, min(j, b) - a
        if first:
            t.text = s[:lo] + new + s[hi:]; first = False
        else:
            t.text = s[:lo] + s[hi:]
        t.set(XMLSPACE, 'preserve')
    return True

def all_pars():
    """Body paragraphs including those in table cells (not the TOC sdt)."""
    out = list(d.paragraphs)
    for tb in d.tables:
        for row in tb.rows:
            for c in row.cells:
                out.extend(c.paragraphs)
    return out

def find(start):
    m = [p for p in d.paragraphs if p.text.strip().startswith(start)]
    if len(m) != 1:
        FAIL.append(f'find({start!r}) matched {len(m)}'); return None
    return m[0]

def rep(section, start, old, new, why):
    p = find(start)
    if p is None:
        return
    before = p.text.strip()
    if not replace_in(p._p, old, new):
        FAIL.append(f'{section}: {old[:50]!r} not in paragraph {start[:40]!r}'); return
    LOG.append((section, why, before, p.text.strip()))

def rep_everywhere(section, old, new, why, expect=None):
    n = 0
    for p in all_pars():
        while old in p.text:
            before = p.text.strip()
            if not replace_in(p._p, old, new):
                break
            n += 1; LOG.append((section, why, before, p.text.strip()))
    if expect is not None and n != expect:
        FAIL.append(f'{section}: {old[:40]!r} replaced {n}x, expected {expect}')
    return n

def delete(section, start, why):
    p = find(start)
    if p is None:
        return
    LOG.append((section, why, p.text.strip(), '(deleted)'))
    p._p.getparent().remove(p._p)

def insert_after(p, text, style):
    new = copy.deepcopy(p._p)
    for child in list(new):
        if child.tag != qn('w:pPr'):
            new.remove(child)
    ppr = new.find(qn('w:pPr'))
    if ppr is not None:
        for tag in ('w:numPr', 'w:pageBreakBefore'):
            e = ppr.find(qn(tag))
            if e is not None:
                ppr.remove(e)
    p._p.addnext(new)
    np_ = Paragraph(new, p._parent); np_.style = d.styles[style]
    np_.add_run(text)
    return np_

def set_lead(section, start, lead, rest, why):
    """Paragraph with a bold run-in lead: keep the bold lead run, replace everything after it."""
    p = find(start)
    if p is None:
        return
    before = p.text.strip()
    runs = list(p._p.iter(qn('w:r')))
    lead_r = runs[0]
    for t in lead_r.iter(qn('w:t')):
        t.text = lead; t.set(XMLSPACE, 'preserve'); break
    second = runs[1]
    ts = list(second.iter(qn('w:t'))); ts[0].text = rest; ts[0].set(XMLSPACE, 'preserve')
    for t in ts[1:]:
        t.text = ''
    for r in runs[2:]:
        r.getparent().remove(r)
    LOG.append((section, why, before, p.text.strip()))


# ================================================================== DECLARATION
# (word count updated at the end, after all edits)

# ================================================================== CHAPTER 1
rep('1.1', 'Large language models (LLMs) offer a different kind',
    'They also hallucinate, vary between runs and can be manipulated by the text they read (Ji et al., 2023; Ouyang et al., 2025; Greshake et al., 2023).',
    'They also hallucinate (Ji et al., 2023), give different answers to the same request (Ouyang et al., 2025) and can be manipulated by the text they read (Greshake et al., 2023).',
    'Each citation placed next to the claim it supports')
rep('1.1', 'Whether an AI reviewer adds value', 'This study investigates it with', 'I investigate it with', 'First person')
rep('1.4', 'The study uses a controlled experimental comparison with mixed-methods',
    'The study uses a controlled experimental comparison with mixed-methods evaluation.',
    'I used a controlled experimental comparison with a mixed-methods evaluation.', 'First person')
rep('1.4', 'After a verification check found',
    'After a verification check found that early AI results had been influenced by scanner output, the AI was run independently and every case was repeated. Outputs were classified against the ground truth, explanation and remediation were scored on a predefined rubric, and the AI’s recommended fixes were applied and re-tested. The contribution is an evidence-based account of how the approaches behave, not a general performance estimate.',
    'When I verified the evidence, I found that early AI results had been influenced by scanner output, so I ran every case again with the AI working independently. I then classified the outputs against the ground truth, scored the explanations and remediation on a predefined rubric, and applied and re-tested the AI’s recommended fixes. The result is an evidence-based account of how the approaches behave on these cases rather than a general performance estimate. My original research proposal is reproduced in Appendix A.',
    'First person; plainer wording; reference to the proposal in Appendix A')
rep('1.5', 'Chapter 5 evaluates the project', 'its validity and the researcher’s own learning.', 'its validity and what I learned.', '"the researcher" → first person')
p = find('Supporting evidence is in the appendices')
if p is not None:
    before = p.text.strip()
    ts = tnodes(p._p)
    ts[0].text = 'Supporting evidence is in the appendices. The complete artefact, including the raw evidence for every run, is publicly available on GitHub ('
    ts[0].set(XMLSPACE, 'preserve')
    hl = p._p.find(qn('w:hyperlink'))
    r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = ').'; r.append(t)
    (hl if hl is not None else p._p).addnext(r)
    LOG.append(('1.5', 'Repository link moved into brackets so the sentence ends cleanly', before, p.text.strip()))
rep('1.6', 'The study involved no human participants and no personal data, and formal', '(Arkko, Cotton and Vegoda, 2010)', '(Arkko et al., 2010)', 'Harvard: three or more authors → et al. (module guidelines)')
rep('1.6', 'Generative AI was both the subject',
    'Generative AI was both the subject of the study and, as permitted by the module guidance, a research assistant for evidence verification, scoring, figure preparation and drafting. Its outputs were checked against primary evidence.',
    'I used generative AI in two ways: as the subject of the study and, as the module guidance permits, as a research assistant for evidence verification, scoring, figure preparation and drafting. I checked its outputs against the primary evidence.',
    'First person (AI-use statement kept, as the module requires)')

# ================================================================== CHAPTER 2
rep('2.2', 'DevSecOps extends DevOps', '(Forsgren, Humble and Kim, 2018)', '(Forsgren et al., 2018)', 'Harvard: et al. for three or more authors')
rep('2.2', 'Two points follow for this study.',
    'Any new reviewer, AI-based or not, should therefore be judged on whether its findings are actionable, not only on how many issues it reports.',
    'Any new reviewer, AI-based or otherwise, should therefore be judged on whether developers can act on its findings as well as on how many issues it reports.', 'Plainer wording')
rep('2.3', 'Research specifically on IaC security', 'A systematic mapping by Rahman, Mahdavi-Hezaveh and Williams (2019)', 'A systematic mapping by Rahman et al. (2019a)', 'Harvard: et al.; 2019a/2019b distinguish two Rahman et al. papers')
rep('2.3', 'Research specifically on IaC security', 'is Rahman, Parnin and Williams (2019).', 'is Rahman et al. (2019b).', 'Harvard: et al. + 2019b')
rep('2.4', 'Static analysis addresses scale.', '(Bridgecrew, 2024; Aqua Security, 2024)', '(Aqua Security, 2024; Bridgecrew, 2024)', 'Same-year citations in alphabetical order')
rep('2.4', 'These are precisely the points',
    'These are precisely the points at which an AI reviewer might add value, and they are where Chapter 4 looks for it.',
    'These are the points at which an AI reviewer might add value, and Chapter 4 examines each of them.', 'Plainer wording')
rep('2.4', 'Most built-in scanner rules encode', '(CIS, 2024; Bridgecrew, 2024)', '(Bridgecrew, 2024; CIS, 2024)', 'Same-year citations in alphabetical order')
rep('2.5', 'The empirical evidence on security tasks', 'Fluent code is therefore not secure code.', 'Code that reads well is therefore not necessarily secure.', 'Plainer wording')
rep('2.5', 'On repair, Pearce et al. (2023)', 'They also show that proposed fixes must be validated, not assumed to work.', 'They also show that a proposed fix has to be tested before it can be trusted.', 'Plainer wording')
rep('2.5', 'Usability studies of AI coding assistants', '(Vaithilingam, Zhang and Glassman, 2022)', '(Vaithilingam et al., 2022)', 'Harvard: et al. for three or more authors')
rep('2.6', 'Remediation is rarely validated.', 'AI remediation is usually assessed by reading it, not by applying it and checking the result.', 'AI remediation is usually judged by reading it; it is rarely applied and checked.', 'Plainer wording')
rep('2.6', 'The realism problem identified by Ding',
    'Controlled cases are therefore suited to explaining how tools behave, not to estimating how often they succeed in practice.',
    'Controlled cases are therefore better suited to explaining how tools behave than to estimating how often they succeed in practice.', 'Plainer wording')
rep('2.7', 'The literature establishes that security review', '(Rahman, Parnin and Williams, 2019;', '(Rahman et al., 2019b;', 'Harvard: et al. + 2019b')
rep_everywhere('Table 2.1', '(Bridgecrew, 2024; Aqua Security, 2024)', '(Aqua Security, 2024; Bridgecrew, 2024)', 'Same-year citations in alphabetical order')

# ================================================================== CHAPTER 3
rep('3.1', 'The design follows the logic of a software engineering experiment',
    'The researcher controls the treatment, a Terraform change with a documented security state. Each reviewer is then applied to the same artefact, and every output is compared with a ground truth fixed before the tools were run.',
    'I controlled the treatment, a Terraform change with a documented security state. Each reviewer was then applied to the same artefact, and I compared every output with a ground truth fixed before the tools were run.',
    '"The researcher" → first person')
rep('3.1', 'The original plan defined three conditions',
    'Between 9 and 12 October 2026 the test cases were run under the workflow default, Condition C, in which normalised Checkov and tfsec findings were inserted into the AI prompt (commit 732083d). Evidence verification on 13 October showed that every AI result from that phase was scanner-informed, so none of it could be reported as AI-only. The workflow was therefore changed (pull request #69, merged on 13 October) so that both scanners and the AI reviewer run in one job without exchanging data. All test cases were then re-run under this independent configuration.',
    'In the first round of runs, the test cases ran under the workflow default, Condition C, in which normalised Checkov and tfsec findings were inserted into the AI prompt (commit 732083d). When I verified the evidence, I found that every AI result from that phase was scanner-informed, so none of it could be reported as AI-only. I therefore changed the workflow (pull request #69) so that both scanners and the AI reviewer run in one job without exchanging data, and then ran every test case under this independent configuration.',
    'Dates removed; first person')
rep('3.3', 'Terraform was chosen because',
    'Terraform was chosen because it is the IaC language used in the Assignment 1 prototype and is widely adopted for multi-cloud provisioning (HashiCorp, 2024). Azure was retained as the target platform for continuity with the existing baseline.',
    'I chose Terraform because it is the IaC language of my Assignment 1 prototype and is widely adopted for multi-cloud provisioning (HashiCorp, 2024). I kept Azure as the target platform for continuity with the existing baseline.',
    'First person')
rep('3.3', 'Checkov and tfsec represent the deterministic', '(Bridgecrew, 2024; Aqua Security, 2024)', '(Aqua Security, 2024; Bridgecrew, 2024)', 'Same-year citations in alphabetical order')
rep('3.3', 'Anthropic Claude was selected in Assignment 1',
    'Anthropic Claude was selected in Assignment 1 as the large language model for the reviewer (Anthropic, 2024). The study evaluates one model and one prompt design. It does not compare models.',
    'I selected Anthropic Claude as the reviewer’s language model in Assignment 1 (Anthropic, 2024). The study evaluates one model and one prompt design and does not compare models.',
    'First person')
rep('3.4', 'The AI reviewer is a Python program', ') that runs only when the pull request', '; Appendix C) that runs only when the pull request', 'Pointer to the code extracts in Appendix C')
rep('3.4', 'Two features of the implementation matter',
    'Second, the prompt was refined on the TC01 baseline before the main experiment, so the reviewer was tuned against part of the test material.',
    'Second, I refined the prompt on the TC01 baseline before the main experiment, so the reviewer was tuned against part of the test material.', 'First person')
rep('3.6', 'The wording used here follows the latest revision',
    'The wording used here follows the latest revision of the test-case evidence page (version 72), which refined several scenario descriptions after the first runs. Where earlier records conflicted, for example over the TC01 baseline or the TC10 severity, the conflict was resolved against the raw run evidence and the decision recorded (Appendix E).',
    'The wording used here follows the latest revision of my test-case evidence page (version 72), which refined several scenario descriptions after the first runs. Where earlier records conflicted, for example over the TC01 baseline or the TC10 severity, I resolved the conflict against the raw run evidence and recorded the decision (Appendix E, Table E.5).',
    'First person; exact table reference')
rep('3.7', 'Explanation and remediation quality were scored',
    'defined in the repository on 8 October, before the independent runs were executed.',
    'defined in the repository (commit c8b2e6f) before the main experimental runs.', 'Date replaced by the commit that proves the order')
rep('3.7', 'The dimensions were fixed in advance.',
    'were written on 13 October after the outputs had been read.', 'were written after I had read the outputs.', 'Date removed; first person')
rep('3.7', 'Scores for all 46 runs were produced',
    'Scores for all 46 runs were produced by an AI research assistant applying these rules, with every score and its written justification recorded alongside the evidence (Appendix E).',
    'I used an AI research assistant to apply these rules to all 46 runs. Every score and its written justification is recorded in the results workbook in the repository and summarised in Appendix E (Table E.3).',
    'First person; accurate location of the justifications')
rep('3.7', 'Remediation effectiveness, the master plan',
    'Remediation effectiveness, the master plan’s demonstrated correction rate, was tested directly.',
    'I tested remediation effectiveness, the master plan’s demonstrated correction rate, directly.', 'First person')
rep('3.8', 'Pull-request comments and job timings were retrieved',
    'Pull-request comments and job timings were retrieved through the GitHub API. Before any result was used, the metadata of all 46 runs was checked to confirm the independent condition, the commit under test and that no scanner findings had been supplied to the AI. Scanner results were also reproduced locally with the same tool versions, and they matched the pipeline outputs in every case.',
    'I retrieved pull-request comments and job timings through the GitHub API. Before using any result, I checked the metadata of all 46 runs to confirm the independent condition, the commit under test and that no scanner findings had been supplied to the AI. I also reproduced the scanner results locally with the same tool versions, and they matched the pipeline outputs in every case.',
    'First person')
rep('3.8', 'Spreadsheet formulas, not hand totals', 'Spreadsheet formulas, not hand totals, produce every reported figure.', 'Every reported figure is calculated by spreadsheet formulas rather than by hand.', 'Plainer wording')
rep('3.9', 'The study involved no human participants and no personal data. All', '(Arkko, Cotton and Vegoda, 2010)', '(Arkko et al., 2010)', 'Harvard: et al.')
rep('3.9', 'Formal ethics approval was not required',
    'In line with the module’s guidance on artificial intelligence, generative AI was used in two ways. It was the object of study, and it also acted as a research assistant for evidence verification, scoring and drafting. All AI-assisted outputs were checked against the primary evidence before use.',
    'In line with the module’s guidance on artificial intelligence, I used generative AI in two ways. It was the object of study, and it also acted as a research assistant for evidence verification, scoring, figure preparation and drafting. I checked all AI-assisted outputs against the primary evidence before using them.',
    'First person; consistent with Section 1.6')

# ================================================================== CHAPTER 4
rep('4.1', 'The figures describe fifteen deliberately', '(Appendix E).', '(Appendices D and E).', 'Run identifiers are in Appendix D')
rep('4.2', 'Both scanners produced identical output', 'Their coverage was therefore almost entirely complementary, not overlapping.', 'Their coverage therefore barely overlapped.', 'Plainer wording')
rep('4.2', 'The most instructive scanner result', 'The most instructive scanner result', 'The most informative scanner result', 'Plainer wording')
rep('4.3', 'However, the two TC12 findings that called', 'They point to a weakness in the test fixtures, not an invented one.', 'They point to a real weakness in the test fixtures.', 'Plainer wording')
rep('4.4', 'In the scanner-informed runs of 9–12 October', 'In the scanner-informed runs of 9–12 October,', 'In the earlier scanner-informed runs,', 'Dates removed')
rep('4.5', 'Figure 4.4 compares detection rates',
    'Figure 4.4 compares detection rates across the fifteen cases.',
    'Figure 4.4 compares detection rates across the fifteen cases, and Table 4.2 summarises the detection and false-alarm measures.', 'Table 4.2 is now referred to in the text')
rep('4.5', 'The size of the gap depends strongly', 'or outside the scanners’ rule sets (TC08, TC09).', 'or outside the scanners’ rule sets (TC09, and TC08 if the partial detection is counted).', 'Accuracy: TC08 is a miss under strict counting')
rep('4.6', 'The most telling false claims', 'The most telling false claims', 'The clearest false claims', 'Plainer wording')
rep('4.6', 'The clearest false claims', '“An argument named”tags” is not expected here”.', '“An argument named ‘tags’ is not expected here” (Appendix F).', 'Broken nested quotation marks fixed; evidence pointer added')
rep('4.7', 'Applying the AI’s own recommendations tested', '(Table 4.3; Figure 4.7)', '(Table 4.3)', 'Figure 4.7 now referenced after Table 4.3, in numbering order')
rep('4.7', 'The fresh AI review accepted nine',
    'The fresh AI review accepted nine of its own fixes. Two of the rest are revealing:',
    'In the fresh AI review, the target condition was no longer raised for nine of the twelve fixes. One of these, TC07, was still blocked: the review raised three other storage settings, one of which it had earlier classified as needing context. Of the other three fixes, two were rejected and one was never reviewed:',
    'Corrected: the count of nine includes TC07, so the old wording listed four cases as "the rest" of three')
set_lead('4.7', 'TC07. The TLS issue', 'TC10.', ' The fix deleted the role definition. Because the runner skips the AI step when a pull request changes no Terraform file, it was never re-reviewed.', 'TC07 moved into the paragraph above; TC10 explained here')
delete('4.7', 'The third case was never re-reviewed', 'Merged into the TC10 bullet above')
rep('4.7', 'These results expose a finding the study did not set out to measure',
    'These results expose a finding the study did not set out to measure: the AI treats natural-language comments in the code as evidence. Four observations support this:',
    'These results also revealed something the study had not set out to measure. The AI treats natural-language comments in the code as evidence, and four observations support this:', 'Plainer wording')
rep('4.8', 'Figure 4.8 shows median step durations', 'Figure 4.8 shows median step durations across the 46 runs:', 'Table 4.4 and Figure 4.8 show the median step durations across the 46 runs.', 'Untitled timing table now captioned and referenced as Table 4.4')
rep('4.9', 'On these cases the AI detected more conditions',
    'Its added detections came from context and semantics, not from better coverage of the scanners’ own rule types.',
    'Its added detections came from reading context and meaning; it did not cover the scanners’ own rule types any better.', 'Plainer wording')
rep('4.9', 'The study did not measure developer understanding directly',
    'The study did not measure developer understanding directly, as no developers took part. The usefulness scores are a proxy judged by the researcher.',
    'RQ3 asks about the likely impact of the AI’s feedback on developer understanding, and I assessed it through the feedback itself. The AI’s comments explained what was wrong, why it mattered and how to fix it, in plain language and specific to the configuration; impact explanation averaged 1.89 and developer usefulness 1.76 out of 2. The scanners’ output was limited to a rule, a short generic description and the affected resource. This points to a positive likely impact on understanding. The main risk is the confident false platform claims described in Section 4.6, because a developer could learn them as fact; explanation accuracy (1.65) was the weakest rubric dimension.',
    'RQ3 answered positively from the evidence; also corrects "judged by the researcher" (the scoring was done with an AI assistant, as Chapter 3 states)')
rep('4.9', 'Taken together, the evidence supports using',
    'Taken together, the evidence supports using the AI reviewer alongside deterministic scanners, not instead of them. It also argues for treating the AI’s verdict as advisory, not as a merge gate.',
    'Taken together, the evidence supports using the AI reviewer alongside deterministic scanners rather than in place of them, with its verdict treated as advice and never as a merge gate.', 'Plainer wording')

# ================================================================== CHAPTER 5
rep('5.1', 'The project built a working pull-request',
    'Every output was classified against ground truth defined before execution. The AI’s explanations were then scored, and its recommended fixes were applied and re-checked.',
    'I classified every output against ground truth defined before execution, scored the AI’s explanations, and applied and re-checked its recommended fixes.', 'First person')
delete('5.2', 'The main shortfall is the human dimension of O4', 'Removed: Objective 4 is fully met as worded (see Table 5.1)')
rep('5.3', 'The TC01 baseline and the prompt were refined after early AI findings',
    'The TC01 baseline and the prompt were refined after early AI findings, so the reviewer was partly tuned on the test material.',
    'I refined the TC01 baseline and the prompt after early AI findings, so the reviewer was partly tuned on the test material.', 'First person')
rep('5.4', 'The most important lesson of this project', 'since 9 October.', 'from the first experimental runs.', 'Date removed')
rep('5.4', 'The project also departed from my PROM05 plan',
    'The project also departed from my PROM05 plan in several ways:',
    'The project also changed from my PROM05 plan in three ways (Appendix G):', 'Clearer; pointer to Appendix G')
set_lead('5.4', 'No developer sessions.', 'From developer sessions to a controlled experiment.',
    ' My PROM05 plan included evaluation sessions with developers. As my thinking developed during Assignment 1, I replaced them with a controlled experiment and measurable criteria, because the quality of the AI’s feedback had to be established before its effect on developers could be meaningful. I therefore answered RQ3 through the feedback itself, which was scored in every run, and the study needed no human participants.',
    'Diplomatic, accurate framing of the scope change (matches the Assignment 1 methodology)')
set_lead('5.4', 'No Azure deployment.', 'No deployment of the test cases.',
    ' The test cases were reviewed as pull requests and never merged or applied, so no deliberately insecure resource was created in Azure. The research questions concern review before deployment, so static analysis of the Terraform was sufficient.',
    'Heading made precise (the baseline environment itself existed in Azure)')
set_lead('5.4', 'Additions.', 'Additions.',
    ' I added repeated runs and remediation testing. Both were responses to weaknesses that verification exposed in the original design.', 'Plainer wording')
rep('5.4', 'I also used an AI assistant to help verify evidence',
    'I also used an AI assistant to help verify evidence, score outputs and draft text.',
    'I also used an AI assistant to help verify evidence, score outputs, prepare figures and draft text.', 'Consistent with the AI-use statement in Sections 1.6 and 3.9')
rep('5.5', 'Professional and legal considerations concern accountability',
    'The findings support keeping the AI advisory, not giving it authority to block or approve merges.',
    'The findings support keeping the AI in an advisory role, without the authority to block or approve merges.', 'Plainer wording')
delete('5.6', 'the absence of developer participants', 'Scope decision now explained once, in Section 5.4')
rep('5.6', 'single-rater scoring;', 'single-rater scoring;', 'single-rater scoring.', 'List punctuation after removing the last item')
delete('5.6', 'Developer study. A small study with developers', 'Duplicate of the future-research point in Section 6.3')

# ================================================================== CHAPTER 6
rep('6.1', 'This was not measured directly.',
    'This was not measured directly. The AI’s comments were consistently more explanatory than the scanners’ output, but whether developers learn from them remains untested.',
    'I assessed this through the quality of the feedback itself. The AI’s comments explained the problem, its impact and the fix in plain language and scored well for usefulness, which points to a positive likely impact on understanding. Its occasional false platform claims are the main risk, because a developer could learn them as fact.',
    'RQ3 conclusion stated positively from the evidence')
rep('6.2', 'Keep deterministic scanners as the enforcement layer',
    'Its findings should prompt human attention, not block or approve merges by themselves.',
    'Its findings should prompt a human to look, and should not block or approve merges by themselves.', 'Plainer wording')

# ================================================================== REFERENCES
rep('References', 'Rahman, A., Mahdavi-Hezaveh, R. and Williams, L. (2019)', '(2019)', '(2019a)', 'Distinguishes the two Rahman et al. (2019) papers')
rep('References', 'Rahman, A., Parnin, C. and Williams, L. (2019)', '(2019)', '(2019b)', 'Distinguishes the two Rahman et al. (2019) papers')
rep('References', 'Ouyang, S., Zhang, J.M.', ', 34(2).', ', 34(2), pp. 1–28. doi:10.1145/3697010.', 'Pages and DOI verified (Crossref)')
rep('References', 'McIntosh, S., Kamei, Y.', 'pp. 2146–2189.', 'pp. 2146–2189. doi:10.1007/s10664-015-9381-9.', 'DOI verified (Crossref)')
rep('References', 'Saavedra, N. and Ferreira, J.F.', '. doi:10.1145/3551349.3556945.', ', pp. 1–12. doi:10.1145/3551349.3556945.', 'Pages verified (Crossref)')
rep('References', 'Li, H., Hao, Y., Zhai, Y. and Qian, Z.', 'Article 111.', 'Article 111. doi:10.1145/3649828.', 'DOI verified (publisher PDF)')
rep('References', 'Tufano, R., Masiero, S.', 'pp. 2291–2302.', 'pp. 2291–2302. doi:10.1145/3510003.3510621.', 'DOI verified (Crossref)')

# ================================================================== APPENDICES
rep('Appendix A', 'Text of the PROM04 research proposal', 'submitted 20 April 2026,', 'submitted for PROM04,', 'Date removed')
rep('Appendix B', 'Version 1.0, 13 Oct 2026. Written and committed', 'Version 1.0, 13 Oct 2026. Written and committed', 'Written and committed (commit cbda52c)', 'Date replaced by the commit reference')
rep('Appendix E', 'Source: results workbook', 'PROM06_TC01-TC12_Verified_Results_Matrix.xlsx', 'PROM06_Results_Workbook.xlsx', 'Workbook renamed in the repository')
rep('Appendix F', 'Source: gh run view 38412955313', '(run by AB, 13 Oct 2026; ANSI colour codes removed)', '(ANSI colour codes removed)', 'Name and date removed (log lines themselves are left verbatim as evidence)')
rep('Appendix F', 'Pages printed to PDF from GitHub', 'Pages printed to PDF from GitHub on 13 October 2026 (18:22–18:29 BST); each page footer', 'Pages printed to PDF from GitHub; each page footer', 'Date removed')
rep_everywhere('Appendix E', 'Agreed (AB, 13 Oct 2026)', 'Adopted', 'Name and date removed', expect=1)
rep('Appendix G', 'The main departures from the PROM05 plan',
    'The main departures from the PROM05 plan were the omission of developer sessions and interviews (and therefore of thematic analysis and ethics review), the use of static review without Azure deployment, the discovery and correction of scanner-informed AI runs, and the addition of repeated runs and remediation validation.',
    'The main changes from the PROM05 plan were a move from developer sessions to a larger controlled experiment, static review without deploying the test cases, the correction of the scanner-informed AI runs, and the addition of repeated runs and remediation validation.',
    'Diplomatic wording, consistent with Section 5.4')

# ------------------------------------------------------------------ tables (cells)
def cell_edit(section, tbl_first_cell, row_key, col, new, why):
    for tb in d.tables:
        if tb.rows[0].cells[0].text.strip() != tbl_first_cell:
            continue
        for row in tb.rows:
            if row.cells[0].text.strip().startswith(row_key):
                c = row.cells[col]; before = c.text.strip()
                ps = c.paragraphs
                ts = tnodes(ps[0]._p)
                if ts:
                    ts[0].text = new; ts[0].set(XMLSPACE, 'preserve')
                    for t in ts[1:]: t.text = ''
                else:
                    ps[0].add_run(new)
                for extra in ps[1:]:
                    extra._p.getparent().remove(extra._p)
                LOG.append((section, why, before, new)); return
    FAIL.append(f'cell {section} {row_key} not found')

cell_edit('Table 5.1', 'Objective', 'O4', 2, 'Achieved', 'Objective 4 is fully met as worded')
G = [('M1', 'Refined in PROM06 Assignment 1'),
     ('M4', 'Pipeline operational; independent configuration introduced through PR #69'),
     ('M5', 'Completed through the controlled experiment: independent runs of all fifteen cases and remediation validation'),
     ('M6', 'Completed (results workbook and evidence register)'),
     ('M7', 'Completed'),
     ('M8', 'This submission')]
for key, new in G:
    cell_edit('Appendix G', 'PROM05 milestone', key, 2, new, 'Dates removed from the plan-versus-actual table')

# ------------------------------------------------------------------ captions for untitled tables
def caption_before_table(tbl_match, text, why):
    cap_tpl = None
    for p in d.paragraphs:
        if p.style.name == 'Table Caption' and p.text.startswith('Table 4.3'):
            cap_tpl = p; break
    for tb in d.tables:
        if tbl_match(tb):
            new = copy.deepcopy(cap_tpl._p)
            for bm in list(new.iter(qn('w:bookmarkStart'))) + list(new.iter(qn('w:bookmarkEnd'))):
                bm.getparent().remove(bm)
            for t in tnodes(new)[1:]:
                t.text = ''
            tnodes(new)[0].text = text
            tb._tbl.addprevious(new)
            LOG.append(('Captions', why, '(no caption)', text)); return
    FAIL.append(f'caption: table not found for {text}')

caption_before_table(lambda tb: tb.rows[0].cells[0].text.strip() == 'Step' and 'Median' in tb.rows[0].cells[1].text,
                     'Table 4.4: Median step durations across the 46 independent runs', 'Timing table had no caption')
caption_before_table(lambda tb: tb.rows[0].cells[0].text.strip() == 'Reviewer',
                     'Table 6.1: Vulnerable cases detected by each reviewer', 'Chapter 6 table had no caption')

# Figure 4.7 referenced after Table 4.3
for tb in d.tables:
    if tb.rows[0].cells[0].text.strip() == 'Case' and 'AI fix applied' in tb.rows[0].cells[1].text:
        tpl = find('Figure 4.4 compares detection rates')
        new = copy.deepcopy(tpl._p)
        for child in list(new):
            if child.tag != qn('w:pPr'):
                new.remove(child)
        tb._tbl.addnext(new)
        np_ = Paragraph(new, tpl._parent); np_.add_run('Figure 4.7 summarises these outcomes.')
        LOG.append(('4.7', 'Figure 4.7 now referenced in the text before it appears', '(none)', 'Figure 4.7 summarises these outcomes.'))

# ------------------------------------------------------------------ Chapter 6: 'Practical challenges' as a heading
p = find('Practical challenges.')
if p is not None:
    before = p.text.strip()
    runs = list(p._p.iter(qn('w:r')))
    for t in runs[0].iter(qn('w:t')):
        t.text = 'Practical challenges'
    for r in runs[1:]:
        r.getparent().remove(r)
    rpr = runs[0].find(qn('w:rPr'))
    if rpr is not None:
        b = rpr.find(qn('w:b'))
        if b is not None: rpr.remove(b)
    tpl_h3 = find('Developer understanding') if False else None
    p.style = d.styles['Heading 3']
    insert_after(p, 'The study found five:', 'Body Text')
    LOG.append(('6.1', 'Run-in label made a sub-heading, like the other three in Section 6.1', before, 'Practical challenges (heading) + "The study found five:"'))

# ------------------------------------------------------------------ layout: chapters on new pages, remove blank-line spacing
body = d.element.body
def is_blank(el):
    if el.tag != qn('w:p'):
        return False
    if ptext(el).strip():
        return False
    for tag in ('w:drawing', 'w:pict', 'w:object', 'w:br', 'w:sectPr', 'w:fldChar', 'w:instrText'):
        if el.find('.//' + qn(tag)) is not None:
            return False
    return True

children = list(body.iterchildren())
# start removing after the declaration's DATE line, stop at Appendix A
start_i = next(i for i, el in enumerate(children) if el.tag == qn('w:p') and ptext(el).strip().startswith('DATE:'))
end_i = next(i for i, el in enumerate(children) if el.tag == qn('w:p') and ptext(el).strip().startswith('Appendix A. Research Proposal'))
removed = 0
for el in children[start_i + 1:end_i]:
    if is_blank(el):
        sty = el.find('.//' + qn('w:pStyle'))
        sname = sty.get(qn('w:val')) if sty is not None else 'Normal'
        if sname in ('CaptionedFigure', 'ImageCaption', 'TableCaption'):
            continue
        body.remove(el); removed += 1
LOG.append(('Layout', f'Removed {removed} empty paragraphs that were used as spacing; page breaks used instead', '', ''))

def pbb(p_el):
    Paragraph(p_el, None).paragraph_format.page_break_before = True

for p in d.paragraphs:
    t = p.text.strip()
    if (p.style.name == 'Heading 1' and t and not t.startswith('Appendix')) or t == 'Abstract':
        pbb(p._p)
sdt = body.find(qn('w:sdt'))
if sdt is not None:
    first_p = sdt.find('.//' + qn('w:p'))
    if first_p is not None:
        pbb(first_p)
LOG.append(('Layout', 'Abstract, contents and every chapter start on a new page', '', ''))

# Appendix A heading levels made consistent with Appendices B–G
n = 0
seen_a = False
for p in d.paragraphs:
    t = p.text.strip()
    if p.style.name == 'Heading 1':
        seen_a = t.startswith('Appendix A')
        continue
    if seen_a and p.style.name == 'Heading 3' and re.match(r'A\.\d ', t):
        p.style = d.styles['Heading 2']; n += 1
    elif seen_a and p.style.name == 'Heading 4' and re.match(r'A\.\d\.\d ', t):
        p.style = d.styles['Heading 3']; n += 1
LOG.append(('Appendix A', f'{n} sub-headings moved up one level to match Appendices B–G', '', ''))

# refresh fields (contents page numbers) when Word opens the file
st = d.settings.element
uf = st.find(qn('w:updateFields'))
if uf is None:
    uf = OxmlElement('w:updateFields'); st.append(uf)
uf.set(qn('w:val'), 'true')

d.save(dst)
json.dump({'log': LOG, 'fail': FAIL}, open(logf, 'w'), ensure_ascii=False, indent=1)
print('changes', len(LOG), 'failures', len(FAIL))
for f in FAIL:
    print('FAIL:', f)
