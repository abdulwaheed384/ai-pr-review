import csv
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

P = "/home/claude/prom06-outputs/PROM06_TC01-TC12_Verified_Results_Matrix.xlsx"
S3 = "/home/claude/prom06-step3"
wb = load_workbook(P)
NEW = ["Metrics", "AI Run Outcomes", "AI Confirmed Findings", "Timing", "Remediation Validation"]
for s in NEW:
    if s in wb.sheetnames:
        del wb[s]

F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F3864")
B = Font(name=F); BB = Font(name=F, bold=True); TI = Font(name=F, bold=True, size=14); IT = Font(name=F, italic=True)
BLUE = Font(name=F, color="0000FF"); GREEN = Font(name=F, color="008000")
thin = Side(style="thin", color="BFBFBF"); BOR = Border(left=thin, right=thin, top=thin, bottom=thin)
WR = Alignment(wrap_text=True, vertical="top"); TOP = Alignment(vertical="top")


def hdr(ws, row, cols, widths=None, start=1):
    for j, c in enumerate(cols, start):
        x = ws.cell(row=row, column=j, value=c); x.font = H; x.fill = HF; x.border = BOR
        x.alignment = Alignment(wrap_text=True, vertical="center")
        if widths:
            ws.column_dimensions[get_column_letter(j)].width = widths[j - start]


def row(ws, r, vals, fonts=None, wrap=()):
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=j, value=v); c.border = BOR
        c.font = (fonts or {}).get(j, B); c.alignment = WR if j in wrap else TOP


# ------------------------------------------------------------------ AI Confirmed Findings
cf = wb.create_sheet("AI Confirmed Findings")
cf["A1"] = "Every CONFIRMED AI finding in the independent condition (46 runs, TC01–TC15) with a manual label against ground truth"; cf["A1"].font = TI
cf["A2"] = ("Source: ai_output.json in each run artefact (TC01–TC12 downloaded by AB on 13 Oct; TC13–TC15 downloaded by AB on 13 Oct, 15:50). "
            "Rows with '-' record runs that produced no confirmed finding. Labels (blue) are the researcher-assistant's judgement; see legend.")
cf["A2"].font = IT; cf["A2"].alignment = WR; cf.merge_cells("A2:I2"); cf.row_dimensions[2].height = 45
cf["A3"] = "Legend: TARGET = the ground-truth condition; EXTRA-VALID = true issue outside the target; DUPLICATE = restates another confirmed finding in the same run; FALSE = not a real issue against ground truth or provider behaviour; NONE = no confirmed finding in the run."
cf["A3"].font = IT; cf["A3"].alignment = WR; cf.merge_cells("A3:I3"); cf.row_dimensions[3].height = 30
LAB = {
    ("TC03", "3", "F-002"): ("DUPLICATE", "Restates F-001 (public RDP) as a SEC-001 policy finding."),
    ("TC07", "1", "F-002"): ("EXTRA-VALID", "Storage account tags lack owner and cost-center (TAG-001); verified in storage.tf."),
    ("TC07", "2", "F-002"): ("EXTRA-VALID", "As run 1."),
    ("TC07", "3", "F-002"): ("EXTRA-VALID", "As run 1."),
    ("TC10", "1", "F-002"): ("FALSE", "azurerm_role_definition has no tags argument; applying the fix fails terraform validate (PR #94)."),
    ("TC10", "3", "F-002"): ("FALSE", "As run 1."),
    ("TC12", "1", "F-001"): ("FALSE", "Public TCP/443 is intended (ground truth TC12); destination-scope concern is not the stated defect."),
    ("TC12", "1", "F-002"): ("FALSE", "Same intended exposure restated against SEC-001."),
    ("TC12", "1", "F-003"): ("FALSE", "azurerm_network_security_rule is not taggable (AI run 3 states this correctly)."),
    ("TC12", "2", "F-001"): ("FALSE", "As TC12 run 1 F-003."),
    ("TC14", "1", "F-001"): ("FALSE", "Ground truth TC14-GT-01: correct outcome is no confirmed finding; wildcard destination is a hardening suggestion, confirmed at LOW."),
}
hdr(cf, 5, ["TC", "Run", "Run ID", "Finding", "Severity", "Title", "Label", "Justification", "Source"], [6, 5, 13, 8, 9, 70, 13, 70, 30])
rows = list(csv.DictReader(open(f"{S3}/confirmed_findings_raw.tsv"), delimiter="\t"))
r = 6
for x in rows:
    if x["finding"] == "-":
        lab, just = "NONE", x["title"].strip("()")
    else:
        lab, just = LAB.get((x["tc"], x["run_no"], x["finding"]), ("TARGET", "Ground-truth condition for this case."))
    row(cf, r, [x["tc"], int(x["run_no"]), x["run_id"], x["finding"], x["severity"], x["title"], lab, just, x["source"]], {7: BLUE}, wrap=(6, 8))
    r += 1
cf_last = r - 1
cf.freeze_panes = "C6"; cf.auto_filter.ref = f"A5:I{cf_last}"

# ------------------------------------------------------------------ AI Run Outcomes
ro = wb.create_sheet("AI Run Outcomes")
ro["A1"] = "AI independent condition — one row per run (46 runs)"; ro["A1"].font = TI
ro["A2"] = ("Counts (black) are formulas over 'AI Confirmed Findings'. Score/decision: TC01–TC12 linked from 'Independent AI Runs' (green); "
            "TC13–TC15 typed from ai_output.json operational_score and verdict in the run artefacts (blue).")
ro["A2"].font = IT; ro["A2"].alignment = WR; ro.merge_cells("A2:L2"); ro.row_dimensions[2].height = 30
hdr(ro, 4, ["TC", "Run", "Run ID", "Case type", "Confirmed findings", "TARGET", "EXTRA-VALID", "DUPLICATE", "FALSE",
            "Target confirmed?", "Operational score", "Decision"], [6, 5, 13, 13, 10, 8, 10, 10, 8, 10, 10, 22])
CT = {"TC01": "Secure", "TC12": "Legitimate", "TC14": "Secure"}
runs = []
seen = set()
for x in rows:
    k = (x["tc"], x["run_no"])
    if k not in seen:
        seen.add(k); runs.append((x["tc"], int(x["run_no"]), x["run_id"]))
import json as _j, glob as _g
T1315 = {rid: (_j.load(open(_g.glob(f"/home/claude/prom06-independent-evidence/artifacts/indep-artifacts/{rid}/*/ai_output.json")[0]))["operational_score"], _j.load(open(_g.glob(f"/home/claude/prom06-independent-evidence/artifacts/indep-artifacts/{rid}/*/ai_output.json")[0]))["verdict"]) for rid in ["38409744211","38409804722","38409830032","38409748586","38409813127","38409838825","38409752049","38409821867","38409847744"]}
_OLD_T1315 = {"38409744211": (70, "APPROVE_WITH_COMMENTS"), "38409804722": (70, "APPROVE_WITH_COMMENTS"), "38409830032": (70, "APPROVE_WITH_COMMENTS"),
         "38409748586": (95, "APPROVE"), "38409813127": (100, "APPROVE"), "38409838825": (100, "APPROVE"),
         "38409752049": (70, "APPROVE_WITH_COMMENTS"), "38409821867": (70, "APPROVE_WITH_COMMENTS"), "38409847744": (70, "APPROVE_WITH_COMMENTS")}
A = f"'AI Confirmed Findings'!$C$6:$C${cf_last}"; G = f"'AI Confirmed Findings'!$G$6:$G${cf_last}"
r = 5
for i, (tc, rn, rid) in enumerate(runs):
    ct = CT.get(tc, "Vulnerable")
    vals = [tc, rn, rid, ct,
            f'=COUNTIFS({A},C{r},{G},"<>NONE")',
            f'=COUNTIFS({A},C{r},{G},"TARGET")', f'=COUNTIFS({A},C{r},{G},"EXTRA-VALID")',
            f'=COUNTIFS({A},C{r},{G},"DUPLICATE")', f'=COUNTIFS({A},C{r},{G},"FALSE")',
            f'=IF(D{r}="Vulnerable",IF(F{r}>0,"Y","N"),"n/a")']
    fonts = {}
    if tc in ("TC13", "TC14", "TC15"):
        vals += list(T1315[rid]); fonts = {11: BLUE, 12: BLUE}
    else:
        src = 2 + i  # Independent AI Runs rows 2..38 are in the same order as the TSV
        vals += [f"='Independent AI Runs'!L{src}", f"='Independent AI Runs'!M{src}"]; fonts = {11: GREEN, 12: GREEN}
    row(ro, r, vals, fonts)
    r += 1
ro_last = r - 1
ro.freeze_panes = "C5"; ro.auto_filter.ref = f"A4:L{ro_last}"

# ------------------------------------------------------------------ Timing
tm = wb.create_sheet("Timing")
tm["A1"] = "Pipeline timing per independent run (GitHub Actions jobs API, step timestamps at 1-second resolution)"; tm["A1"].font = TI
tm["A2"] = ("All three tools run sequentially in one job and all PR comments are posted together at the end, so time-to-feedback is the same for every tool; "
            "tool cost is shown as step duration. AI duration includes the Anthropic API call. PR→comments is only meaningful for pull_request-event runs.")
tm["A2"].font = IT; tm["A2"].alignment = WR; tm.merge_cells("A2:L2"); tm.row_dimensions[2].height = 30
hdr(tm, 4, ["Run ID", "Branch", "Event", "Job start (UTC)", "Job (s)", "Install Checkov (s)", "Run Checkov (s)", "Run tfsec (s)",
            "AI review (s)", "AI step result", "PR created (UTC)", "PR → comments posted (s)"], [13, 12, 16, 21, 8, 10, 9, 9, 9, 10, 21, 12])
ev = {}
for line in open("/home/claude/prom06-independent-evidence/independent-runs.tsv").read().splitlines()[1:]:
    p = line.split("\t"); ev[p[2]] = p[1]
for line in open(f"{S3}/tc13-15-runs.tsv").read().splitlines():
    p = line.split("\t"); ev[p[0]] = p[2]
tr = list(csv.DictReader(open(f"{S3}/step-timings.tsv"), delimiter="\t"))
tr.sort(key=lambda d: (d["branch"], d["job_start"]))
r = 5
from datetime import datetime
for d in tr:
    e = ev.get(d["run_id"], "")
    if e == "pull_request" and d["comments_posted"]:
        lat = (datetime.fromisoformat(d["comments_posted"].replace("Z", "+00:00")) - datetime.fromisoformat(d["pr_created"].replace("Z", "+00:00"))).total_seconds()
    else:
        lat = "n/a"
    row(tm, r, [d["run_id"], d["branch"], e, d["job_start"].replace("T", " ").replace("Z", ""), float(d["job_s"]),
                float(d["checkov_install_s"]), float(d["checkov_s"]), float(d["tfsec_s"]), float(d["ai_s"]), d["ai_concl"],
                d["pr_created"].replace("T", " ").replace("Z", ""), lat])
    r += 1
tm_last = r - 1
s = tm_last + 2
tm.cell(row=s, column=1, value="Summary (formulas)").font = TI
hdr(tm, s + 1, ["Statistic", "", "", "", "Job (s)", "Install Checkov (s)", "Run Checkov (s)", "Run tfsec (s)", "AI review (s)", "", "", "PR → comments (s)"])
for k, (lab, fn) in enumerate((("Median", "MEDIAN"), ("Minimum", "MIN"), ("Maximum", "MAX"), ("Mean", "AVERAGE")), s + 2):
    tm.cell(row=k, column=1, value=lab).font = BB
    for col in "EFGHIL":
        c = tm[f"{col}{k}"]; c.value = f"={fn}({col}5:{col}{tm_last})"; c.number_format = "0.0"
tm.cell(row=s + 6, column=1, value="Runs").font = BB; tm[f"E{s+6}"] = f"=COUNT(E5:E{tm_last})"
tm.cell(row=s + 7, column=1, value="AI step share of job time (median AI ÷ median job)").font = BB
tm[f"E{s+7}"] = f"=I{s+2}/E{s+2}"; tm[f"E{s+7}"].number_format = "0%"
tm.freeze_panes = "B5"

# ------------------------------------------------------------------ Remediation Validation
rv = wb.create_sheet("Remediation Validation")
rv["A1"] = "Remediation validation (DCR) — does applying the AI's own recommendation produce the ground-truth secure state?"; rv["A1"].font = TI
rv["A2"] = ("Method (13 Oct 2026): for each case where the AI confirmed the target, the recommendation from the first run that confirmed it was applied to the "
            "unchanged fixture branch on a new rem-tc branch, with no other edits. Each branch was opened as a PR (not merged, not applied) so the INDEPENDENT workflow ran "
            "terraform fmt/validate, Checkov 3.3.1, tfsec 1.28.14 and a fresh AI review once. Where the AI gave a placeholder source range, the RFC 5737 documentation "
            "address 203.0.113.10/32 was used (no real addresses). Fixes were implemented by the research assistant (Claude); interpretation choices are recorded. "
            "Variant 'b' rows change only the now-stale fixture comment, to isolate its effect on the AI re-review. One AI run per branch: no variability estimate.")
rv["A2"].font = IT; rv["A2"].alignment = WR; rv.merge_cells("A2:P2"); rv.row_dimensions[2].height = 75
cols = ["TC", "Variant", "AI recommendation source", "Fix applied", "Interpretation / note", "Branch", "Commit", "PR", "Run ID",
        "fmt + validate", "Checkov target check (before → after)", "tfsec (before → after)", "AI re-review (independent, 1 run)",
        "Target resolved vs ground truth?", "Scanner confirms fix?", "AI accepts fix?"]
hdr(rv, 4, cols, [6, 7, 22, 34, 34, 13, 9, 6, 13, 9, 26, 24, 44, 12, 11, 10])
REM = [
 ("TC02", "a", "Run 38404584887 F-001", "source_address_prefix 0.0.0.0/0 → 203.0.113.10/32", "AI example: 'jump-host IP'; documentation address used. Fixture comment still says 'Deliberately unrestricted SSH'.", "rem-tc02", "386c2ac", 85, "38412911309", "Pass", "CKV_AZURE_10 passed before and after (blind to 0.0.0.0/0)", "AVD-AZU-0047 + 0050 → 0", "70 APPROVE_WITH_COMMENTS — HIGH: source is an RFC 5737 TEST-NET address; cites the stale comment", "Yes", "Yes (tfsec)", "No"),
 ("TC02", "b", "As variant a", "As a + comment updated to describe the restricted rule", "Isolates stale-comment effect.", "rem-tc02b", "fd061b8", 98, "38413482687", "Pass", "Passed (7/0)", "0 → 0", "95 APPROVE — LOW: wildcard destination (same note as TC14 run 1)", "Yes", "Yes (tfsec)", "Yes"),
 ("TC03", "a", "Run 38404589228 F-001", "source_address_prefix 0.0.0.0/0 → 203.0.113.10/32", "AI: 'controlled IP range … or Azure Bastion'; first option applied. Stale comment retained.", "rem-tc03", "08e5734", 86, "38412915964", "Pass", "CKV_AZURE_9 passed before and after (blind)", "AVD-AZU-0048 + 0047 → 0", "70 APPROVE_WITH_COMMENTS — HIGH: RDP from external source address; quotes 'Deliberately unrestricted RDP' comment", "Yes", "Yes (tfsec)", "No"),
 ("TC03", "b", "As variant a", "As a + comment updated", "Isolates stale-comment effect.", "rem-tc03b", "ed0ab21", 99, "38413403063", "Pass", "Passed (7/0)", "0 → 0", "100 APPROVE — no confirmed findings", "Yes", "Yes (tfsec)", "Yes"),
 ("TC04", "a", "Run 38404592498 F-001", "source_address_prefix 0.0.0.0/0 → 203.0.113.10/32", "AI: restrict to test-harness range. No HTTPS alternative was recommended. Stale comment retained.", "rem-tc04", "8d02bac", 87, "38412921766", "Pass", "CKV_AZURE_160 passed before and after (blind)", "AVD-AZU-0047 → 0", "95 APPROVE — LOW: comment contradicts the configured source prefix", "Yes", "Yes (tfsec)", "Yes"),
 ("TC04", "b", "As variant a", "As a + comment updated", "Isolates stale-comment effect.", "rem-tc04b", "6e7a619", 100, "38413413620", "Pass", "Passed (7/0)", "0 → 0", "85 APPROVE — MEDIUM: HTTP is cleartext, use HTTPS (an issue the AI did not raise when reviewing the original fixture)", "Yes", "Yes (tfsec)", "Yes"),
 ("TC05", "a", "Run 38404594354 F-001", "Add azurerm_subnet_network_security_group_association", "Applied as given.", "rem-tc05", "952f4dc", 88, "38412929801", "Pass", "CKV2_AZURE_31 failed → passed", "0 → 0", "100 APPROVE — no confirmed findings", "Yes", "Yes (Checkov)", "Yes"),
 ("TC06", "a", "Run 38404596950 F-001", "Remove dns_servers (Azure-provided DNS)", "First of two AI options (other: 168.63.129.16).", "rem-tc06", "fb59f8a", 89, "38412933192", "Pass", "CKV_AZURE_183 + 182 failed → passed", "0 → 0", "100 APPROVE — no confirmed findings", "Yes", "Yes (Checkov)", "Yes"),
 ("TC07", "a", "Run 38404600206 F-001", "min_tls_version TLS1_0 → TLS1_2", "Target finding only; other AI recommendations not applied.", "rem-tc07", "6d60bce", 90, "38412938552", "Pass", "CKV_AZURE_44 failed → passed (10 unrelated storage checks still fail)", "AVD-AZU-0011 → 0", "40 DO_NOT_MERGE — TLS no longer raised; 3 other confirmed: tags (true), public network access HIGH (was NEEDS_CONTEXT), no diagnostic setting", "Yes", "Yes (both)", "Target yes; gate blocks"),
 ("TC08", "a", "Run 38405218950 F-001", "Add azurerm_monitor_diagnostic_setting for the NSG with both categories", "Only run that confirmed the target; code example applied as given.", "rem-tc08", "324f05c", 91, "38412942226", "Pass", "7/0 before and after (blind)", "0 → 0", "100 APPROVE — no confirmed findings", "Yes", "n/a (blind)", "Yes"),
 ("TC09", "a", "Run 38404607277 F-001", "Remove the locals block containing the dummy secret", "AI option 'remove the locals block'. Value remains in git history (not recommended for purge in this run).", "rem-tc09", "a3323f5", 92, "38412949170", "Pass", "7/0 before and after (blind; secrets framework not run)", "0 → 0", "100 APPROVE — no confirmed findings", "Yes (code); history not addressed", "n/a (blind)", "Yes"),
 ("TC10", "a", "Run 38404612132 F-001", "Delete the role definition", "First AI option. Second option ('minimum set of actions') could not be applied: the AI did not specify the actions.", "rem-tc10", "46226ce", 93, "38412950863", "Pass", "CKV_AZURE_39 failed → passed", "0 → 0", "Not invoked: PR has no .tf diff against main after deletion", "Yes", "Yes (Checkov)", "n/a"),
 ("TC10", "tags", "Run 38404612132 F-002 (FALSE finding)", "Add tags = local.common_tags to azurerm_role_definition", "Applied verbatim to test the false tag finding.", "rem-tc10tags", "7c47e50", 94, "38412955313", "FAIL — Unsupported argument: \"tags\" not expected on azurerm_role_definition (later steps skipped)", "Skipped", "Skipped", "Skipped", "n/a — fix is invalid Terraform", "n/a", "n/a"),
 ("TC11", "a", "Run 38404613836 F-001", "actions → the two documented read actions; assignable_scopes → resource group", "Matches tc11_required_actions exactly.", "rem-tc11", "2e55a55", 95, "38412957421", "Pass", "CKV_AZURE_39 failed → passed (8/0)", "0 → 0", "100 APPROVE — no confirmed findings", "Yes", "Yes (Checkov)", "Yes"),
 ("TC13", "a", "Run 38409744211 F-001", "role Contributor → Reader; scope subscription → resource group", "Applied as given.", "rem-tc13", "3748898", 96, "38412963217", "Pass", "7/0 before and after (blind)", "0 → 0", "100 APPROVE — no confirmed findings", "Yes", "n/a (blind)", "Yes"),
 ("TC15", "a", "Run 38409752049 F-001", "Remove default; add sensitive = true", "Rotation and history purge are operational steps outside code; not tested.", "rem-tc15", "30a6cf7", 97, "38412968077", "Pass", "7/0 before and after (blind)", "0 → 0", "95 APPROVE — LOW: no validation block / no confirmed secret store", "Yes (code)", "n/a (blind)", "Yes (LOW note)"),
]
r = 5
for x in REM:
    row(rv, r, list(x), wrap=(3, 4, 5, 11, 12, 13, 14))
    rv.row_dimensions[r].height = 48
    r += 1
rv_last = r - 1
rv.freeze_panes = "C5"

s = rv_last + 2
rv.cell(row=s, column=1, value="Summary (formulas; primary variant 'a' unless stated)").font = TI
lab = [
    ("Cases with an AI remediation tested (variant a)", f'=COUNTIFS(B5:B{rv_last},"a")'),
    ("Target resolved vs ground truth (variant a; 'Yes…' incl. qualified)", f'=COUNTIFS(B5:B{rv_last},"a",N5:N{rv_last},"Yes*")'),
    ("Passed terraform fmt + validate (variant a)", f'=COUNTIFS(B5:B{rv_last},"a",J5:J{rv_last},"Pass")'),
    ("Fix confirmed by a scanner that had detected the target", f'=COUNTIFS(B5:B{rv_last},"a",O5:O{rv_last},"Yes*")'),
    ("Cases where neither scanner could confirm (blind)", f'=COUNTIFS(B5:B{rv_last},"a",O5:O{rv_last},"n/a (blind)")'),
    ("AI re-review accepted its own fix (variant a, 'Yes…')", f'=COUNTIFS(B5:B{rv_last},"a",P5:P{rv_last},"Yes*")+COUNTIFS(B5:B{rv_last},"a",P5:P{rv_last},"Target yes*")'),
    ("AI re-review rejected its own fix with a HIGH finding (variant a)", f'=COUNTIFS(B5:B{rv_last},"a",P5:P{rv_last},"No")'),
    ("Variant b (stale comment corrected, TC02–TC04): AI accepted the fix", f'=COUNTIFS(B5:B{rv_last},"b",P5:P{rv_last},"Yes*")'),
    ("AI remediation for a false finding that broke validation", f'=COUNTIFS(J5:J{rv_last},"FAIL*")'),
]
for k, (l, f) in enumerate(lab, s + 1):
    rv.cell(row=k, column=1, value=l).font = BB; rv.merge_cells(start_row=k, start_column=1, end_row=k, end_column=5)
    rv.cell(row=k, column=6, value=f).font = B
rv.cell(row=s + len(lab) + 2, column=1, value=("rem-tc10tags validate error (log supplied by AB, 13 Oct 2026; archived at step3/rem-logs/38412955313-terraform-validate.txt): "
        "Error: Unsupported argument on tc10_iam.tf line 14 — An argument named \"tags\" is not expected here.")).font = IT

# ------------------------------------------------------------------ Metrics
m = wb.create_sheet("Metrics")
m["A1"] = "Step 3 metrics — detection, false alarms, run consistency, timing and remediation"; m["A1"].font = TI
m["A2"] = ("Case-level classes for TC01–TC12 are linked from 'Verified Matrix' (green). TC13–TC15 are typed (blue) from local scans and the nine run artefacts (metadata verified: INDEPENDENT, no scanner input). "
           " All rates are descriptive for constructed cases (n = 15), not population estimates.")
m["A2"].font = IT; m["A2"].alignment = WR; m.merge_cells("A2:H2"); m.row_dimensions[2].height = 45
hdr(m, 4, ["TC", "Case type", "Checkov", "tfsec", "Scanner union", "AI independent", "Status", "Note"], [8, 13, 11, 11, 13, 14, 12, 60])
for i in range(12):
    r = 5 + i; vm = 2 + i
    m.cell(row=r, column=1, value=f"='Verified Matrix'!A{vm}").font = GREEN
    m.cell(row=r, column=2, value=f"='Verified Matrix'!D{vm}").font = GREEN
    m.cell(row=r, column=3, value=f"='Verified Matrix'!F{vm}").font = GREEN
    m.cell(row=r, column=4, value=f"='Verified Matrix'!H{vm}").font = GREEN
    m.cell(row=r, column=6, value=f"='Verified Matrix'!J{vm}").font = GREEN
    m.cell(row=r, column=7, value="Verified").font = B
extra = [("TC13", "Vulnerable", "FN", "FN", "TP", "Contributor at subscription vs documented Reader-on-RG need. Checkov 7/0, tfsec 0; AI HIGH 3/3."),
         ("TC14", "Secure", "TN", "TN", "TN", "Restricted SSH (remediated TC02). AI: LOW wildcard-destination finding confirmed in run 1; runs 2–3 clean with NEEDS_CONTEXT bastion notes (hidden from PR comment)."),
         ("TC15", "Vulnerable", "FN", "FN", "TP", "Hard-coded password in variable default. Checkov 7/0 (secrets framework not run), tfsec 0; AI HIGH 3/3.")]
for k, (tc, ct, c, t, a, note) in enumerate(extra):
    r = 17 + k
    for j, v in ((1, tc), (2, ct), (3, c), (4, t), (6, a)):
        m.cell(row=r, column=j, value=v).font = BLUE
    m.cell(row=r, column=7, value="Verified").font = B
    m.cell(row=r, column=8, value=note).font = B
m["H5"] = "Secure baseline."; m["H12"] = "Raised 4/4 runs, confirmed 1/4 (PARTIAL)."; m["H13"] = "Only the AI detected it (scanners FN)."
m["H15"] = "Construct not isolated: CKV_AZURE_39 fires on the wildcard alone."; m["H16"] = "Report both readings (decision D2)."
for r in range(5, 20):
    m.cell(row=r, column=5, value=f'=IF(OR(C{r}="TP",D{r}="TP"),"TP",IF(B{r}="Vulnerable","FN",IF(OR(C{r}="CONTEXT",D{r}="CONTEXT"),"CONTEXT",IF(OR(C{r}="FP",D{r}="FP"),"FP","TN"))))')
    for j in range(1, 9):
        m.cell(row=r, column=j).border = BOR
        if m.cell(row=r, column=j).font != GREEN and m.cell(row=r, column=j).font != BLUE:
            pass
    m.cell(row=r, column=8).alignment = WR

# case-level metric blocks
def block(top, title, last):
    m.cell(row=top, column=1, value=title).font = TI
    hdr(m, top + 1, ["Measure", "", "Checkov", "tfsec", "Scanner union", "AI independent", "", "Definition"])
    R_ = lambda col: f"{col}$5:{col}${last}"
    ct = R_("$B")
    defs = [
        ("Vulnerable cases", lambda c: f'=COUNTIF({ct},"Vulnerable")', "Denominator for detection rate"),
        ("Detected (TP)", lambda c: f'=COUNTIFS({ct},"Vulnerable",{R_(c)},"TP")', "One detection per case regardless of rule count"),
        ("Partial (raised, not consistently confirmed)", lambda c: f'=COUNTIFS({ct},"Vulnerable",{R_(c)},"PARTIAL")', "AI TC08"),
        ("Missed (FN)", lambda c: f'=COUNTIFS({ct},"Vulnerable",{R_(c)},"FN")', ""),
        ("Detection rate, strict (TP ÷ vulnerable)", None, "PARTIAL counted as miss"),
        ("Detection rate, lenient ((TP + PARTIAL) ÷ vulnerable)", None, "PARTIAL counted as hit"),
        ("Non-vulnerable cases (secure + legitimate)", lambda c: f'=COUNTIF({ct},"<>Vulnerable")', "TC01, TC12 (+TC14)"),
        ("Clean on non-vulnerable (TN)", lambda c: f'=COUNTIFS({ct},"<>Vulnerable",{R_(c)},"TN")', ""),
        ("Flagged non-vulnerable (FP or CONTEXT)", lambda c: f'=COUNTIFS({ct},"<>Vulnerable",{R_(c)},"FP")+COUNTIFS({ct},"<>Vulnerable",{R_(c)},"CONTEXT")', "TC12 counted here for tfsec and AI"),
    ]
    for k, (lab, fn, d) in enumerate(defs):
        r = top + 2 + k
        m.cell(row=r, column=1, value=lab).font = BB
        m.cell(row=r, column=8, value=d).font = IT
        for col in "CDEF":
            c = m[f"{col}{r}"]
            if lab.startswith("Detection rate, strict"):
                c.value = f"=IFERROR({col}{top+3}/{col}{top+2},0)"; c.number_format = "0%"
            elif lab.startswith("Detection rate, lenient"):
                c.value = f"=IFERROR(({col}{top+3}+{col}{top+4})/{col}{top+2},0)"; c.number_format = "0%"
            else:
                c.value = fn(col)
    return top + 2 + len(defs)

end1 = block(22, "A. Case level — TC01–TC12 (verified)", 16)
end2 = block(end1 + 2, "B. Case level — TC01–TC15", 19)

# complementarity
t = end2 + 2
m.cell(row=t, column=1, value="C. Complementarity — vulnerable cases, TC01–TC15").font = TI
V = "($B$5:$B$19=\"Vulnerable\")"
comp = [
    ("Detected by Checkov OR tfsec", f'=SUMPRODUCT({V}*($E$5:$E$19="TP"))'),
    ("Detected by AI (TP)", f'=SUMPRODUCT({V}*($F$5:$F$19="TP"))'),
    ("Detected by AI but by neither scanner", f'=SUMPRODUCT({V}*($F$5:$F$19="TP")*($E$5:$E$19<>"TP"))'),
    ("Detected by a scanner but not by AI (TP)", f'=SUMPRODUCT({V}*($E$5:$E$19="TP")*($F$5:$F$19<>"TP"))'),
    ("Detected by scanners OR AI (any tool)", f'=SUMPRODUCT({V}*((($E$5:$E$19="TP")+($F$5:$F$19="TP"))>0))'),
    ("Missed by every tool (counting AI PARTIAL as a miss)", f'=SUMPRODUCT({V}*($E$5:$E$19<>"TP")*($F$5:$F$19<>"TP"))'),
]
for k, (lab, f) in enumerate(comp, t + 1):
    m.cell(row=k, column=1, value=lab).font = BB; m.cell(row=k, column=3, value=f)
m.cell(row=t + 1, column=8, value="Scanner union").font = IT
m.cell(row=t + 3, column=8, value="TC09, TC13, TC15 (TC08 is PARTIAL)").font = IT
m.cell(row=t + 6, column=8, value="TC08 (AI PARTIAL; scanners FN)").font = IT

# run-level AI
u = t + len(comp) + 2
m.cell(row=u, column=1, value="D. AI run level — 46 independent runs").font = TI
RO = lambda col: f"'AI Run Outcomes'!${col}$5:${col}${ro_last}"
CFL = f"'AI Confirmed Findings'!$G$6:$G${cf_last}"
runm = [
    ("Vulnerable-case runs", f'=COUNTIF({RO("D")},"Vulnerable")', "0"),
    ("…target CONFIRMED", f'=COUNTIF({RO("J")},"Y")', "0"),
    ("Run-level detection rate", f"=C{u+2}/C{u+1}", "0%"),
    ("Non-vulnerable runs (TC01, TC12, TC14)", f'=COUNTIF({RO("D")},"<>Vulnerable")', "0"),
    ("…with ≥1 confirmed finding (false alarm)", f'=COUNTIFS({RO("D")},"<>Vulnerable",{RO("E")},">0")', "0"),
    ("Run-level false-alarm rate", f"=C{u+5}/C{u+4}", "0%"),
    ("Non-vulnerable runs that blocked the gate (DO_NOT_MERGE)", f'=COUNTIFS({RO("D")},"<>Vulnerable",{RO("L")},"DO_NOT_MERGE")', "0"),
    ("All confirmed findings", f'=COUNTIF({CFL},"<>NONE")', "0"),
    ("…TARGET", f'=COUNTIF({CFL},"TARGET")', "0"),
    ("…EXTRA-VALID", f'=COUNTIF({CFL},"EXTRA-VALID")', "0"),
    ("…DUPLICATE", f'=COUNTIF({CFL},"DUPLICATE")', "0"),
    ("…FALSE", f'=COUNTIF({CFL},"FALSE")', "0"),
    ("Precision of confirmed findings ((TARGET + EXTRA-VALID) ÷ all)", f"=(C{u+9}+C{u+10})/C{u+8}", "0%"),
    ("Runs with ≥1 FALSE confirmed finding", f'=COUNTIF({RO("I")},">0")', "0"),
    ("Runs that failed output validation", f'=COUNTIF(\'AI Confirmed Findings\'!$F$6:$F${cf_last},"*failed validation*")', "0"),
]
for k, (lab, f, nf) in enumerate(runm, u + 1):
    m.cell(row=k, column=1, value=lab).font = BB
    c = m.cell(row=k, column=3, value=f); c.number_format = nf
m.cell(row=u + 1, column=8, value="TC02–TC11 (31) + TC13, TC15 (6)").font = IT
m.cell(row=u + 2, column=8, value="Misses: TC08 runs 1–3 (FALSE_POSITIVE, NEEDS_CONTEXT ×2)").font = IT

# timing + remediation links
v = u + len(runm) + 2
m.cell(row=v, column=1, value="E. Timing (median seconds, 46 runs) — linked from 'Timing'").font = TI
ts = tm_last + 4  # median row in Timing
for k, (lab, col) in enumerate((("Whole job", "E"), ("Run Checkov (excl. install)", "G"), ("Install Checkov", "F"), ("Run tfsec", "H"),
                                 ("AI review step", "I"), ("PR opened → all comments posted (PR runs)", "L")), v + 1):
    m.cell(row=k, column=1, value=lab).font = BB
    c = m.cell(row=k, column=3, value=f"=Timing!{col}{ts}"); c.font = GREEN; c.number_format = "0.0"
w = v + 8
m.cell(row=w, column=1, value="F. Remediation validation (DCR) — linked from 'Remediation Validation'").font = TI
rs = rv_last + 3
for k, lab in enumerate(["AI remediations tested", "Target resolved vs ground truth", "Passed fmt + validate", "Confirmed by a scanner that had detected the target",
                         "Scanner blind (could not confirm)", "AI re-review accepted its own fix", "AI rejected its own fix with HIGH (stale comment)",
                         "Variant b (comment corrected): AI accepted, of 3", "Remediation of a false finding broke validation"], w + 1):
    m.cell(row=k, column=1, value=lab).font = BB
    m.cell(row=k, column=3, value=f"='Remediation Validation'!F{rs + (k - w - 1)}").font = GREEN

# fix stale note on Summary
sm = wb["Summary"]
for rr in range(1, 30):
    if sm.cell(row=rr, column=6).value == "Pending researcher decision":
        sm.cell(row=rr, column=6).value = "Agreed 13 Oct: report both readings (D2)"

# order sheets: Metrics after Summary; detail sheets after Inter-rater Check
order = ["README", "Verified Matrix", "Summary", "Metrics", "Independent AI Runs", "AI Run Outcomes", "AI Confirmed Findings", "Scanner Results",
         "Hybrid Runs (Condition C)", "Timing", "Remediation Validation", "Decisions & Caveats", "Scoring Rules", "AI Quality Scoring",
         "Inter-rater Check", "Draft Matrix (original)"]
wb._sheets = [wb[n] for n in order]
wb.save(P)
print("saved", cf_last, ro_last, tm_last, rv_last)
