import sys, docx
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docxcompose.composer import Composer
master_p, app_p, out = sys.argv[1:4]
app = docx.Document(app_p)
def no_num(p):
    pPr = p._p.get_or_add_pPr(); n = pPr.find(qn('w:numPr'))
    if n is None: n = OxmlElement('w:numPr'); pPr.insert(0, n)
    for c in list(n): n.remove(c)
    il = OxmlElement('w:ilvl'); il.set(qn('w:val'), '0'); n.append(il)
    ni = OxmlElement('w:numId'); ni.set(qn('w:val'), '0'); n.append(ni)
for p in app.paragraphs:
    if p.style.name.startswith('Heading'):
        no_num(p)
        if p.style.name == 'Heading 1': p.paragraph_format.page_break_before = True
# tables: borders, single spacing, small font, header bold
def borders(t):
    tblPr = t._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for e in ('top','left','bottom','right','insideH','insideV'):
        x = OxmlElement(f'w:{e}'); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'999999'); b.append(x)
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None: tblPr.remove(old)
    tblPr.append(b)
for t in app.tables:
    borders(t)
    for i, row in enumerate(t.rows):
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.size = Pt(9 if len(t.columns) > 5 else 10)
                    if i == 0: r.bold = True
# code blocks single spaced
for p in app.paragraphs:
    if p.style.name in ('Source Code',): p.paragraph_format.line_spacing = 1.0
app.save('app_fmt.docx')
m = docx.Document(master_p)
c = Composer(m); c.append(docx.Document('app_fmt.docx'))
c.save(out); print('composed')
