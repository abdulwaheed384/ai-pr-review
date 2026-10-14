import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt, numpy as np, matplotlib.ticker
from openpyxl import load_workbook
wb=load_workbook("/home/claude/prom06-outputs/PROM06_TC01-TC12_Verified_Results_Matrix.xlsx",data_only=True); met=wb["Metrics"]
rows=[r for r in range(20,60) if met.cell(row=r,column=1).value=="Detection rate, strict (TP ÷ vulnerable)"]
b=[met.cell(row=rows[1],column=c).value for c in (3,4,5,6)]
print(b)
SURF="#fcfcfb"; INK="#0b0b0b"; INK2="#52514e"; GRID="#e4e3df"; S1="#2a78d6"; S2="#eb6834"
plt.rcParams.update({"font.family":"DejaVu Sans","font.size":10,"axes.edgecolor":INK2,"axes.labelcolor":INK2,"xtick.color":INK2,"ytick.color":INK2,"figure.facecolor":SURF,"axes.facecolor":SURF,"savefig.facecolor":SURF})
tools=["Checkov","tfsec","Checkov ∪ tfsec","AI independent"]
fig,ax=plt.subplots(figsize=(7.2,3.6))
x=np.arange(4); bars=ax.bar(x,b,0.55,color=[S1,S1,S1,S2])
for r_,v in zip(bars,b): ax.text(r_.get_x()+r_.get_width()/2,v+0.02,f"{v:.0%}",ha="center",va="bottom",fontsize=9,color=INK)
ax.set_xticks(x,tools); ax.set_ylim(0,1.1); ax.yaxis.set_major_formatter(matplotlib.ticker.PercentFormatter(1.0))
ax.set_ylabel("Vulnerable cases detected (strict)")
for s in("top","right"): ax.spines[s].set_visible(False)
ax.grid(axis="y",color=GRID,linewidth=0.8); ax.set_axisbelow(True)
ax.text(0,-0.24,"12 vulnerable cases in TC01–TC15. Strict: AI PARTIAL (TC08) counted as a miss; the lenient AI rate is 100%.",transform=ax.transAxes,fontsize=7.5,color=INK2)
fig.savefig("/home/claude/prom06-outputs/figures/fig_detection_rates_15.png",dpi=300,bbox_inches="tight")
