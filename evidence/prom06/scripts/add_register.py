from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

P = "/home/claude/prom06-outputs/PROM06_TC01-TC12_Verified_Results_Matrix.xlsx"
wb = load_workbook(P)
for s in ("Evidence Register", "Figures & Tables", "Inter-rater Check"):
    if s in wb.sheetnames:
        del wb[s]
F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F3864")
B = Font(name=F); BB = Font(name=F, bold=True); TI = Font(name=F, bold=True, size=14); IT = Font(name=F, italic=True)
thin = Side(style="thin", color="BFBFBF"); BOR = Border(left=thin, right=thin, top=thin, bottom=thin)
WR = Alignment(wrap_text=True, vertical="top")
EB = "branch prom06-evidence: evidence/prom06/"


def sheet(name, title, note, cols, widths, rows, status_col):
    ws = wb.create_sheet(name)
    ws["A1"] = title; ws["A1"].font = TI
    ws["A2"] = note; ws["A2"].font = IT; ws["A2"].alignment = WR
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(cols)); ws.row_dimensions[2].height = 45
    for j, (c, w) in enumerate(zip(cols, widths), 1):
        x = ws.cell(row=4, column=j, value=c); x.font = H; x.fill = HF; x.border = BOR; x.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = w
    for i, r in enumerate(rows, 5):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v); c.font = B; c.alignment = WR; c.border = BOR
    last = 4 + len(rows)
    rng = f"{get_column_letter(status_col)}5:{get_column_letter(status_col)}{last}"
    for txt, col in (("Verified", "C6EFCE"), ("Provisional", "FFEB9C"), ("Pending AB", "FFC7CE"), ("To produce", "FFEB9C"), ("Ready", "C6EFCE")):
        ws.conditional_formatting.add(rng, CellIsRule(operator="beginsWith" if False else "equal", formula=[f'"{txt}"'], fill=PatternFill("solid", fgColor=col)))
    ws.freeze_panes = "B5"; ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{last}"
    return ws, last


R = [
 # ID, chapter, claim, evidence type, location, status, caveat for the write-up
 ("E01", "Ch3", "The study compares three conditions: A scanners only, B AI only, C hybrid.", "Planning document", "Project doc 08 Master Plan §3", "Verified", "Executed as: scanners + AI INDEPENDENT (A and B in one job) and earlier Condition C runs."),
 ("E02", "Ch3", "From 9 Oct 2026 12:47 the workflow default (Condition C) injected normalised scanner findings into the AI prompt via build_scanner_context().", "Repository history", "ai-pr-review commit 732083d", "Verified", "All hybrid results TC01–TC12 come from this period."),
 ("E03", "Ch3", "PR #69 (merged 13 Oct 09:59, merge commit 3a79020) made the workflow INDEPENDENT; metadata records scanner_findings_supplied_to_ai = false.", "Repository + run metadata", "PR #69; " + EB + "independent-runs/<run_id>/experiment-metadata.json", "Verified", ""),
 ("E04", "Ch3", "Every independent run (46) executed with INDEPENDENT condition and no scanner input to the AI.", "Run metadata", EB + "independent-runs/ (46 runs, TC13–TC15 artefacts supplied by AB 13 Oct 15:50)", "Verified", "Metadata lacks the model name (hard-coded in runner)."),
 ("E05", "Ch3", "Tool configuration: Checkov 3.3.1 (--framework terraform, secrets framework not run), tfsec 1.28.14, AI claude-sonnet-4-6 via API, max_tokens 3000, no temperature set, at most 3 findings, classification policy 2.2.", "Code", ".github/workflows/terraform.yml; ai/review_runner.py; metadata", "Verified", "Model name is hard-coded, not recorded in run metadata."),
 ("E06", "Ch3", "The AI receives the changed .tf files plus every terraform/*.tf file, ai/prompt.txt, ai/policy.md and README.md.", "Code", "ai/review_runner.py get_terraform_review_context() (lines 28–64)", "Verified", ""),
 ("E07", "Ch3", "Only CONFIRMED AI findings appear in the PR comment and the operational score.", "Code", "ai/review_runner.py format_comment(), calculate_operational_score(); ai/scoring.json", "Verified", ""),
 ("E08", "Ch3", "Ground truth for TC01–TC12 was documented before tool execution; TC13–TC15 ground truth was committed before their fixtures.", "Planning doc + git order", "Ground-Truth Matrix / Confluence v72; commit cbda52c precedes 1e87e00, 5a48d73, 049f9c9", "Verified", "TC01–TC12 wording revised in Confluence (v72); use v72."),
 ("E09", "Ch3", "Independent fixtures TC02–TC12 are byte-identical to the original test-case Terraform; TC01 differs by one comment.", "git diff", "branches indep-tc01…12 vs original TC commits", "Verified", ""),
 ("E10", "Ch3", "Each case ran three times on the same commit (TC08 four times, PR run failed).", "Run list", EB + "independent-runs.tsv; step3/tc13-15-runs.tsv", "Verified", ""),
 ("E11", "Ch4", "Scanner output was identical across runs and matched a local reproduction with the same versions.", "Raw JSON", EB + "local-scans/TC01–TC15; independent-runs/*/checkov-results.json, tfsec-results.json", "Verified", ""),
 ("E12", "Ch4", "Detection (TC01–TC12, 10 vulnerable): Checkov 5, tfsec 4, scanner union 8, AI independent 9 + 1 partial.", "Workbook", "Metrics block A; Verified Matrix", "Verified", "Descriptive only; constructed cases."),
 ("E13", "Ch4", "Detection (TC01–TC15, 12 vulnerable): Checkov 5, tfsec 4, union 8, AI 11 + 1 partial; AI alone found TC09, TC13, TC15; no scanner-only detection.", "Workbook", "Metrics blocks B–C", "Verified", "Descriptive only; constructed cases."),
 ("E14", "Ch4", "Checkov did not flag SSH/RDP/HTTP open to 0.0.0.0/0 (TC02–TC04) but CKV_AZURE_10 fails when the source is '*'.", "Local probe", EB + "step3/probe-tc02-star/", "Verified", "Rule value-matching blind spot, not a configuration error."),
 ("E15", "Ch4", "TC08: the AI raised the missing NSG diagnostic setting in all four runs but classified it FALSE_POSITIVE (run failed validation), NEEDS_CONTEXT twice and CONFIRMED once.", "AI output", EB + "independent-runs/38404605125 (ai_raw_response.txt), 38405066857, 38405174844, 38405218950", "Verified", "Classification instability, not a detection miss."),
 ("E16", "Ch4", "TC12 (legitimate HTTPS) varied across identical runs: blocked (50), LOW tag finding (95), clean (100).", "AI output", EB + "independent-runs/38404619147, 38405103514, 38405210047", "Verified", ""),
 ("E17", "Ch4", "AI run level: target confirmed in 34/37 vulnerable-case runs; false alarm in 3/9 non-vulnerable runs; one gate block.", "Workbook", "Metrics block D; AI Run Outcomes", "Verified", ""),
 ("E18", "Ch4", "45 confirmed AI findings: 34 target, 3 extra-valid, 1 duplicate, 7 false → 82% precision.", "Workbook", "AI Confirmed Findings (labels by research assistant)", "Verified", "Labels are the assistant's judgement against ground truth; AB to spot-check the 7 FALSE labels."),
 ("E19", "Ch4", "The AI claimed TAG-001 applies to azurerm_role_definition and azurerm_network_security_rule; applying its tag fix failed terraform validate.", "AI output + CI run", "TC10 runs 1, 3; TC12 runs 1, 2; PR #94 run 38412955313; " + EB + "step3/rem-logs/38412955313-terraform-validate.txt", "Verified", "Validate error: Unsupported argument — An argument named \"tags\" is not expected here (tc10_iam.tf line 14)."),
 ("E20", "Ch4", "AI quality (37 vulnerable-case runs): mean 0–2 scores — see AI Quality Scoring summary (revised 13 Oct); 46 runs scored in total.", "Workbook", "AI Quality Scoring; Scoring Rules", "Verified", "Single rubric application by AI assistant; no independent second rater (limitation)."),
 ("E21", "Ch4", "9 of 46 runs mapped findings to an inapplicable policy, mostly SEC-001; ai/policy.md has no encryption, IAM or secrets policy.", "Workbook + code", "AI Quality Scoring col P; ai/policy.md", "Verified", "Input asymmetry: policy.md is given to the AI only."),
 ("E22", "Ch4", "Median timing: job 77 s, AI step 31.5 s, Checkov 6 s (+19.5 s install), tfsec 1 s, PR opened → comments 79 s.", "GitHub jobs API", EB + "step3/jobs/; Timing sheet", "Verified", "1-second resolution; sequential steps."),
 ("E23", "Ch4", "Applying each AI fix resolved the target in 12/12 cases and passed fmt/validate; 8 confirmed by the scanner that had flagged the issue.", "CI runs + local scans", "PRs #85–#97; " + EB + "step3/rem-*; Remediation Validation", "Verified", "Fixes implemented by the research assistant; one run each."),
 ("E24", "Ch4", "With stale 'Deliberately unrestricted' comments, the AI rejected its own TC02/TC03 fixes (HIGH); with only the comment corrected it accepted them.", "CI runs", "PR #85 vs #98; #86 vs #99", "Verified", "Single run per variant."),
 ("E25", "Ch4", "The AI quotes in-code 'Ground truth:' comments (TC12 run 3); TC11 fixture contains a similar comment.", "AI output + fixture", "38405210047 ai_output.json F-002; indep-tc11 tc11_requirements.tf", "Verified", "Cueing threat; links to prompt-injection literature."),
 ("E26", "Ch4", "A PR that only deletes .tf files is not reviewed by the AI.", "Code + CI run", "review_runner.py (is_file filter); PR #93 AI comment", "Verified", "Deletion-only changes bypass AI review."),
 ("E27", "Ch4", "Hybrid (Condition C) AI detected 10/10 vulnerable cases while citing scanner rule IDs; TC01 hybrid result is unverified.", "PR comments", "Hybrid Runs (Condition C); PRs #46, #55–#68", "Verified", "Not comparable with independent results: scanner-informed."),
 ("E28", "Ch4", "Hybrid severities were higher than independent (SSH/RDP CRITICAL vs HIGH); TC10 hybrid gave 2×CRITICAL then 1×HIGH on the same commit.", "PR comments", "PRs #56, #57, #66 vs #71, #72, #79", "Verified", ""),
 ("E29", "Ch5", "The TC01 baseline and prompt were changed after early AI findings (PR #46 → #47–#54).", "Repository history", "PRs #46–#55", "Verified", "Circularity threat."),
 ("E30", "Ch5", "TC07 is not single-variable: Checkov fails 10 unrelated storage checks.", "Raw JSON", EB + "local-scans/TC07/checkov.json", "Verified", ""),
 ("E31", "Ch5", "TC11 construct not isolated: CKV_AZURE_39 fires on the wildcard alone.", "Raw JSON", EB + "local-scans/TC11", "Verified", ""),
 ("E32", "Ch5", "The AI quality scorer was an AI assistant that had read the outputs before writing operational rules.", "Process record", "Scoring Rules sheet", "Verified", "Disclose; no inter-rater check performed (AB decision, 13 Oct)."),
 ("E33", "Ch5", "The Confluence evidence page calls the Condition C AI results 'independent' and labels TC11 as Condition B.", "Confluence", "Page 131215 v72", "Verified", "Page left unchanged at AB's decision; explain in Ch5 if relevant."),
 ("E34", "Ch5", "Submission deadline 26 Oct 2026 23:59.", "Canvas export", "Project doc 11", "Verified", "From Canvas export; AB chose not to re-check live Canvas (13 Oct)."),
 ("E36", "Ch4", "Explanations overstated technical facts in some runs: Contributor described as able to change role assignments (TC13 runs 2–3, contradicted by Microsoft's role definition); an unreferenced local said to be present in state now (TC09 run 2).", "AI output + vendor docs", EB + "independent-runs/38409804722, 38409830032, 38405075984; Microsoft Learn Azure built-in roles (Contributor NotActions)", "Verified", "Conditional state-exposure claims (TC09 runs 1, 3; TC15 run 2) are consistent with HashiCorp guidance and were not penalised."),
 ("E37", "Ch4", "With its secrets framework enabled, Checkov flags the TC09 string as high-entropy (CKV_SECRET_6) but not the readable password in TC15.", "Local probe", EB + "step3/probe-secrets/", "Verified", "Probe outside the experimental configuration (pipeline runs --framework terraform only)."),
 ("E35", "Ch3", "No test PR was merged and the experimental workflow contains no terraform apply step.", "GitHub PR state + workflow", "Test PRs #56–#68 and #70–#100 unmerged (PR #69, the workflow change, was merged); .github/workflows/terraform.yml", "Verified", "A Terraform Cloud check also ran on PRs (speculative plan, to confirm); close PRs unmerged after sign-off."),
]
ws1, l1 = sheet("Evidence Register", "Evidence register — every factual claim the dissertation will make, traced to its evidence",
                "Status: Verified = checked against primary evidence; Provisional = depends on TC13–TC15 artefacts or AB's inter-rater check; Pending AB = needs an action from AB. "
                "Write-up rule: a claim may appear in the dissertation only if it is here, and Provisional/Pending claims must be re-checked first.",
                ["ID", "Chapter", "Claim", "Evidence type", "Location", "Status", "Caveat for the write-up"], [6, 7, 70, 16, 50, 12, 45], R, 6)

FT = [
 ("Figure", "F3.1", "Ch3", "Experimental pipeline and conditions", "figures/fig_architecture.png (Graphviz source .dot)", "Ready", "Shows Condition C injection path."),
 ("Figure", "F4.1", "Ch4", "Detection outcome by case and tool (TC01–TC15)", "figures/fig_detection_matrix.png", "Ready", ""),
 ("Figure", "F4.2", "Ch4", "Detection rate by tool, two case sets", "figures/fig_detection_rates.png", "Ready", ""),
 ("Figure", "F4.3", "Ch4", "AI explanation/remediation quality by rubric dimension", "figures/fig_ai_quality.png", "Ready", ""),
 ("Figure", "F4.4", "Ch4", "Median pipeline timings", "figures/fig_timing.png", "Ready", ""),
 ("Figure", "F4.5", "Ch4", "Remediation validation outcomes", "figures/fig_remediation.png", "Ready", ""),
 ("Screenshot", "S1", "Ch4", "PR #71 (TC02): Checkov clean vs tfsec and AI findings", "GitHub PR page", "Pending AB", ""),
 ("Screenshot", "S2", "Ch4", "PR #78 (TC09): AI-only detection of an embedded secret", "GitHub PR page", "Pending AB", ""),
 ("Screenshot", "S3", "Ch4", "PR #77 (TC08): failed run / validation rejection", "GitHub Actions run 38404605125", "Pending AB", ""),
 ("Screenshot", "S4", "Ch4", "PR #81 (TC12): block on legitimate HTTPS (run 1)", "GitHub PR page", "Pending AB", ""),
 ("Screenshot", "S5", "Ch4", "PRs #82 and #84 (TC13, TC15): AI-only detections", "GitHub PR pages", "Pending AB", ""),
 ("Screenshot", "S6", "Ch4", "PR #85 vs #98: stale comment changes the AI verdict", "GitHub PR pages", "Pending AB", ""),
 ("Screenshot", "S7", "Ch4", "PR #94: AI tag fix fails terraform validate", "GitHub Actions run 38412955313 (log text already archived)", "Pending AB", "Optional: log text can be shown as a code listing instead."),
 ("Screenshot", "S8", "Ch3", "Run summary page showing uploaded artefacts", "Any independent run", "Pending AB", ""),
 ("Table", "T3.1", "Ch3", "Test cases TC01–TC15 with ground truth and purpose", "Confluence v72; ground-truth/TC13-TC15", "To produce", ""),
 ("Table", "T3.2", "Ch3", "Tool versions and configuration", "E05, E06", "To produce", ""),
 ("Table", "T3.3", "Ch3", "Classification codes and normalisation rules", "Verified Matrix README; Decisions & Caveats", "To produce", ""),
 ("Table", "T3.4", "Ch3", "Remediation-quality rubric with operational rules", "Scoring Rules", "To produce", ""),
 ("Table", "T4.1", "Ch4", "Case-level detection metrics", "Metrics A–C", "To produce", ""),
 ("Table", "T4.2", "Ch4", "AI run-level metrics and confirmed-finding precision", "Metrics D", "To produce", ""),
 ("Table", "T4.3", "Ch4", "False confirmed AI findings (7) with explanation", "AI Confirmed Findings", "To produce", ""),
 ("Table", "T4.4", "Ch4", "Remediation validation by case", "Remediation Validation", "To produce", ""),
 ("Table", "T4.5", "Ch4", "Independent vs hybrid AI (TC01–TC12)", "Hybrid Runs; Verified Matrix", "To produce", ""),
 ("Table", "T5.1", "Ch5", "Threats to validity and mitigations", "E29–E33, Decisions & Caveats", "To produce", ""),
 ("Appendix", "App A", "—", "Evidence index (branch, paths, run IDs, PRs)", "This register + evidence README", "To produce", ""),
 ("Appendix", "App B", "—", "Ground truth TC01–TC15", "Confluence v72; ground-truth file", "To produce", ""),
 ("Appendix", "App C", "—", "AI prompt and policy (prompt.txt, policy.md)", "Repository", "To produce", ""),
 ("Appendix", "App D", "—", "Full per-run AI scores and confirmed findings", "AI Quality Scoring; AI Confirmed Findings", "To produce", ""),
 ("Appendix", "App E", "—", "Project management record", "Project doc 07 (PROM05 plan) + actual timeline", "To produce", ""),
]
ws2, l2 = sheet("Figures & Tables", "Planned figures, screenshots, tables and appendices (provisional numbering)",
                "Figures F* are rendered at 300 dpi from this workbook by scripts in the evidence branch. Screenshots S* must be captured by AB from GitHub (the assistant cannot screenshot web pages). Excluded from the word count per the brief.",
                ["Kind", "No.", "Chapter", "Content", "Source", "Status", "Note"], [11, 7, 8, 55, 50, 12, 35], FT, 6)

order = wb.sheetnames
order.remove("Evidence Register"); order.remove("Figures & Tables")
i = order.index("Metrics") + 1
order[i:i] = ["Evidence Register", "Figures & Tables"]
wb._sheets = [wb[n] for n in order]
wb.save(P)
print(l1, l2, wb.sheetnames)
