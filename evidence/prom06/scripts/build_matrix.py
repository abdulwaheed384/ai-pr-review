import json,glob,os,collections
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ART="/home/claude/prom06-independent-evidence/artifacts/indep-artifacts"
LOC="/home/claude/prom06-local-scans"
OUT="/home/claude/prom06-outputs/PROM06_TC01-TC12_Verified_Results_Matrix.xlsx"
F="Arial"
H=Font(name=F,bold=True,color="FFFFFF"); HF=PatternFill("solid",fgColor="1F3864")
B=Font(name=F); BB=Font(name=F,bold=True); TI=Font(name=F,bold=True,size=14)
thin=Side(style="thin",color="BFBFBF"); BOR=Border(left=thin,right=thin,top=thin,bottom=thin)
WR=Alignment(wrap_text=True,vertical="top")
FILL={"TP":"E2EFDA","TN":"E2EFDA","FN":"FCE4D6","FP":"F8CBAD","PARTIAL":"FFF2CC","CONTEXT":"FFF2CC","UNVERIFIED":"EDEDED"}

shamap={"92ab884":"TC01","3eda6f7":"TC02","27789d9":"TC03","cbc6855":"TC04","14486bd":"TC05","eb51ee9":"TC06","b6e93b7":"TC07","c86e4d5":"TC08","277a730":"TC09","bd8a61f":"TC10","3dd9a04":"TC11","9cee1fe":"TC12"}
pr={"TC%02d"%i:69+i for i in range(1,13)}

def header(ws,row,cols,widths):
    for j,(c,w) in enumerate(zip(cols,widths),1):
        x=ws.cell(row=row,column=j,value=c); x.font=H; x.fill=HF; x.alignment=WR; x.border=BOR
        ws.column_dimensions[get_column_letter(j)].width=w
def put(ws,row,vals,bold=False):
    for j,v in enumerate(vals,1):
        x=ws.cell(row=row,column=j,value=v); x.font=BB if bold else B; x.alignment=WR; x.border=BOR

wb=Workbook()
# ---------- README
ws=wb.active; ws.title="README"
lines=[("PROM06 TC01–TC12 Verified Results Matrix",TI),
("Built 13 Oct 2026 from: PR bot comments and run artifacts (PRs #46–#68 Condition C; PRs #70–#81 independent), experiment-metadata.json for 37 independent runs, local reproduction with Checkov 3.3.1 and tfsec 1.28.14, Ground-Truth Matrix v1.0 and Confluence PROM06-TC-Evidences v72 (13 Oct 10:47).",B),
("Evidence hierarchy used: raw run artifacts > PR bot comments > Confluence/screenshots > draft matrix.",B),
("Classification codes (target condition only):",BB),
("TP = target condition correctly reported   FN = target condition not reported   TN = no finding on a case with no intended weakness",B),
("FP = finding that treats an acceptable configuration as a defect   PARTIAL = reported inconsistently across runs   CONTEXT = context-sensitive, researcher decision needed   UNVERIFIED = evidence conflict, not usable",B),
("AI (independent) = post-PR #69 workflow, scanner_findings_supplied_to_ai = false in every run's metadata, 3 successful runs per case (TC08: 3 successful + 1 failed). AI (hybrid) = earlier Condition C runs whose prompt contained normalised Checkov/tfsec findings.",B),
("Severity labels are each tool's own scale and are not equivalent across tools (Checkov OSS JSON has no severity).",B),
("Sheets: Verified Matrix · Summary (formula-driven counts) · Independent AI Runs · Scanner Results · Hybrid Runs (Condition C) · Decisions & Caveats · Draft Matrix (original).",B),
("Yellow cells on 'Decisions & Caveats' are items for AB to confirm before the matrix is frozen.",B)]
for i,(t,f) in enumerate(lines,1):
    c=ws.cell(row=i,column=1,value=t); c.font=f; c.alignment=Alignment(wrap_text=True)
ws.column_dimensions["A"].width=140

# ---------- Independent runs data
runs=[]
for m in glob.glob(ART+"/*/*/experiment-metadata.json"):
    d=os.path.dirname(m); md=json.load(open(m)); tc=shamap[md["experiment_commit_sha"][:7]]
    ai=json.load(open(d+"/ai_output.json")) if os.path.exists(d+"/ai_output.json") else None
    runs.append((tc,md,ai))
runs.sort(key=lambda r:(r[0],r[1]["started_at_utc"]))

# ---------- Verified matrix
vm=wb.create_sheet("Verified Matrix")
cols=["Test case","Scenario (Confluence v72)","Ground truth (summary)","Case type","Checkov 3.3.1 observed","Checkov class.","tfsec 1.28.14 observed","tfsec class.","AI independent (3 runs)","AI indep. class.","AI hybrid / Condition C (earlier)","AI hybrid class.","Extra / unrelated findings","Evidence references","Caveats"]
header(vm,1,cols,[8,24,34,12,30,10,30,10,40,11,34,11,38,30,40])
M=[
("TC01","Secure baseline","No intentional vulnerability (false-positive baseline).","Secure",
 "7 passed / 0 failed","TN","0 findings","TN",
 "0 confirmed in 3/3 runs (100/100). Only FALSE_POSITIVE/INFORMATIONAL notes in the artifact.","TN",
 "Conflicting: PR #46 (older fixture) 2 confirmed MEDIUM (TAG-001, LOG-001), 55/100; Confluence 'clean' screenshot from PR #55 cannot be tied to its own run.","UNVERIFIED",
 "None confirmed.","PR #70, 92ab884, runs 38404582550 / 38405003190 / 38405111880; hybrid PR #46, #55","Baseline evolved after PR #46 AI findings (tags, logging, deny rules) and prompt was adjusted — disclose as validity threat. TC01 rerun branch has a comment-only edit."),
("TC02","Unrestricted SSH Access","Inbound TCP/22 from 0.0.0.0/0 — genuine management-plane exposure.","Vulnerable",
 "7/0 — CKV_AZURE_10 evaluated and PASSED","FN","AVD-AZU-0047 + AVD-AZU-0050 (CRITICAL)","TP",
 "HIGH SSH finding in 3/3 runs (70/100).","TP",
 "CRITICAL SSH, cited tfsec IDs (PR #56).","TP",
 "tfsec: 2 rules for one condition (normalise to 1).","PR #71, 3eda6f7; hybrid PR #56","Checkov passes because '0.0.0.0/0' is not matched as internet; CKV_AZURE_10 fails when the source is '*' (local probe)."),
("TC03","Unrestricted RDP Access","Inbound TCP/3389 from 0.0.0.0/0 — genuine management-plane exposure.","Vulnerable",
 "7/0 — CKV_AZURE_9 evaluated and PASSED","FN","AVD-AZU-0048 + AVD-AZU-0047 (CRITICAL)","TP",
 "HIGH RDP finding in 3/3 runs; run 3 added a duplicate SEC-001 policy finding (40/100).","TP",
 "CRITICAL RDP, cited tfsec IDs (PR #57).","TP",
 "AI run 3: duplicate SEC-001 finding for the same condition; tfsec: 2 rules, 1 condition.","PR #72, 27789d9; hybrid PR #57","Same Checkov source-matching behaviour as TC02 (inferred; probe run on TC02 only)."),
("TC04","Public HTTP Ingress","Inbound TCP/80 from 0.0.0.0/0 — intended TC04 weakness (Confluence v72).","Vulnerable",
 "7/0 — CKV_AZURE_160 evaluated and PASSED","FN","AVD-AZU-0047 (CRITICAL, generic public ingress)","TP",
 "HIGH HTTP finding in 3/3 runs (70/100).","TP",
 "HIGH HTTP, cited tfsec ID (PR #58).","TP",
 "—","PR #73, cbc6855; hybrid PR #58","GT v1.0 described TC04 as context-dependent; Confluence v72 fixes it as an intended weakness. tfsec rule is not port-specific."),
("TC05","Missing NSG Association","NSG exists but is not associated with the subnet — control not effective.","Vulnerable",
 "CKV2_AZURE_31 (6 passed / 1 failed)","TP","0 findings","FN",
 "NET-001 finding in 3/3 runs (HIGH, HIGH, MEDIUM).","TP",
 "MEDIUM NET-001, cited Checkov ID (PR #59).","TP",
 "AI: NEEDS_CONTEXT note in every run about the stale 'Custom resolver' DNS comment (fixture artefact).","PR #74, 14486bd; hybrid PR #59","Severity varied across identical runs."),
("TC06","Public DNS Configuration","VNet uses public resolver 8.8.8.8 contrary to NET-002 / CKV_AZURE_183.","Vulnerable",
 "CKV_AZURE_183 + CKV_AZURE_182 (5/2)","TP","0 findings","FN",
 "MEDIUM NET-002 finding in 3/3 runs (85/100).","TP",
 "MEDIUM NET-002, cited Checkov IDs (PR #61).","TP",
 "Checkov CKV_AZURE_182 (≥2 DNS endpoints) is related but not the target condition.","PR #75, eb51ee9; hybrid PR #61","Ground truth depends on the project baseline (NET-002); public DNS is not inherently insecure."),
("TC07","Weak TLS Configuration","Storage account min_tls_version = TLS1_0 below TLS 1.2 minimum.","Vulnerable",
 "CKV_AZURE_44 among 11 failed (11/11)","TP","AVD-AZU-0011 (CRITICAL)","TP",
 "HIGH TLS finding in 3/3 runs + LOW TAG-001 (65/100).","TP",
 "CRITICAL TLS + 2 extra findings (PR #63).","TP",
 "Checkov: 10 unrelated storage checks fail. AI: valid LOW TAG-001 (storage is taggable) + NEEDS_CONTEXT public access.","PR #76, b6e93b7; hybrid PR #63","Fixture not single-variable (new storage account, hard-coded RG, non-standard tags); Confluence scenario said TC12 HTTPS service would be reused."),
("TC08","Missing Diagnostic/Audit Logging","NSG diagnostic setting removed; logs not forwarded to Log Analytics (LOG-001).","Vulnerable",
 "7/0","FN","0 findings","FN",
 "Raised in all 4 runs but classified inconsistently: FALSE_POSITIVE (output rejected, run failed), NEEDS_CONTEXT ×2, CONFIRMED ×1.","PARTIAL",
 "MEDIUM LOG-001 confirmed (PR #64).","TP",
 "—","PR #77, c86e4d5, runs 38404605125 (failed) / 38405066857 / 38405174844 / 38405218950; hybrid PR #64","Only the AI is given LOG-001; failed run's summary wrongly claimed the absence was intentional."),
("TC09","Plaintext Credentials (dummy)","Dummy credential-shaped string hard-coded in a Terraform local; value is non-sensitive test data.","Vulnerable",
 "7/0 (Checkov run with --framework terraform; secrets framework not run)","FN","0 findings","FN",
 "MEDIUM hard-coded secret finding in 3/3 runs (85/100).","TP",
 "MEDIUM finding (PR #65).","TP",
 "—","PR #78, 277a730; hybrid PR #65","Checkov non-detection reflects pipeline configuration as much as tool capability. Explanation accuracy (state claims) to be scored in step 2."),
("TC10","Excessive Azure IAM Permissions","Custom role actions = ['*'] at subscription scope exceeds minimum.","Vulnerable",
 "CKV_AZURE_39 (7/1)","TP","0 findings","FN",
 "HIGH wildcard-role finding in 3/3 runs; 2/3 runs added LOW TAG-001.","TP",
 "Run 1: 2 × CRITICAL (40/100); rerun same commit: 1 × HIGH (70/100) (PR #66).","TP",
 "AI TAG-001 on azurerm_role_definition (resource has no tags) — AI false positive in 2/3 runs.","PR #79, bd8a61f; hybrid PR #66 (2 attempts)","Matrix 'critical' reflected hybrid run 1 only."),
("TC11","Multi-file, context-dependent weakness","Wildcard role vs read-only actions documented in tc11_requirements.tf.","Vulnerable",
 "CKV_AZURE_39 (7/1)","TP","0 findings","FN",
 "HIGH finding in 3/3 runs; all 3 cite tc11_requirements.tf.","TP",
 "HIGH finding ×2, cited Checkov (PR #67).","TP",
 "—","PR #80, 3dd9a04; hybrid PR #67 (2 attempts)","Construct validity: the wildcard alone triggers detection, so this does not isolate multi-file reasoning. Confluence heading 'AI-only (B)' is wrong — evidence is Condition C."),
("TC12","Legitimate Public HTTPS Service","Public TCP/443 is intentional; exposure alone is not a vulnerability.","Legitimate",
 "11 passed / 0 failed","TN","AVD-AZU-0047 (CRITICAL) on the intended rule","CONTEXT",
 "Varied: run 1 blocked (HIGH + MEDIUM SEC-001 + LOW TAG-001, 50/100); run 2 LOW TAG-001 only (95); run 3 none (100), said SEC-001 justified and TAG-001 n/a.","CONTEXT",
 "MEDIUM SEC-001 scope concern (PR #68).","CONTEXT",
 "AI TAG-001 on azurerm_network_security_rule (resource has no tags) — AI false positive in 2/3 runs.","PR #81, 9cee1fe; hybrid PR #68","Decide whether the shared-NSG/wildcard-destination scope concern counts as a separate weakness; until then tfsec and AI are CONTEXT."),
]
for i,r in enumerate(M,2):
    put(vm,i,r)
    for col in (6,8,10,12):
        c=vm.cell(row=i,column=col); c.fill=PatternFill("solid",fgColor=FILL[c.value]); c.font=BB
vm.freeze_panes="B2"
n=len(M)+1
vm.auto_filter.ref=f"A1:O{n}"

# ---------- Summary with formulas
sm=wb.create_sheet("Summary")
sm["A1"]="Descriptive counts (formula-driven from 'Verified Matrix')"; sm["A1"].font=TI
header(sm,3,["Measure","Checkov","tfsec","AI independent","AI hybrid (Cond. C)","Definition / denominator"],[46,12,12,16,18,60])
R=f"'Verified Matrix'!$D$2:$D${n}"
def cnt(col,code,typ="Vulnerable"): return f"=COUNTIFS({R},\"{typ}\",'Verified Matrix'!${col}$2:${col}${n},\"{code}\")"
cls={"Checkov":"F","tfsec":"H","AIi":"J","AIh":"L"}
rows=[
("Vulnerable cases (denominator)",*[f"=COUNTIF({R},\"Vulnerable\")"]*4,"TC02–TC11"),
("Target condition detected (TP)",*[cnt(c,"TP") for c in cls.values()],"Normalised: one detection per case regardless of rule count"),
("Target condition missed (FN)",*[cnt(c,"FN") for c in cls.values()],""),
("Inconsistent across runs (PARTIAL)",*[cnt(c,"PARTIAL") for c in cls.values()],"AI only: TC08"),
("Detection rate (TP / vulnerable cases)",*[f"=IFERROR({x}5/{x}4,0)" for x in "BCDE"],"Descriptive only; n = 10 constructed cases, not a population estimate"),
("Secure baseline TC01 — clean (TN)",*[cnt(c,"TN","Secure") for c in cls.values()],"1 case"),
("Secure baseline TC01 — unverified",*[cnt(c,"UNVERIFIED","Secure") for c in cls.values()],"Hybrid TC01 evidence conflict"),
("Legitimate TC12 — no finding (TN)",*[cnt(c,"TN","Legitimate") for c in cls.values()],"1 case"),
("Legitimate TC12 — context-sensitive (CONTEXT)",*[cnt(c,"CONTEXT","Legitimate") for c in cls.values()],"Pending researcher decision"),
]
for i,r in enumerate(rows,4):
    put(sm,i,r)
    if i==8:
        for j in range(2,6): sm.cell(row=i,column=j).number_format="0%"
sm["A15"]="Union and overlap (vulnerable cases)"; sm["A15"].font=BB
F_=f"'Verified Matrix'!$F$2:$F${n}"; Hh=f"'Verified Matrix'!$H$2:$H${n}"; J=f"'Verified Matrix'!$J$2:$J${n}"
extra=[
("Detected by Checkov OR tfsec",f"=SUMPRODUCT(({R}=\"Vulnerable\")*((({F_}=\"TP\")+({Hh}=\"TP\"))>0))","Scanner union"),
("Detected by AI independent (TP) but by neither scanner",f"=SUMPRODUCT(({R}=\"Vulnerable\")*({J}=\"TP\")*({F_}<>\"TP\")*({Hh}<>\"TP\"))","TC09 (TC08 is PARTIAL)"),
("Detected by a scanner but not by AI independent (TP)",f"=SUMPRODUCT(({R}=\"Vulnerable\")*(((({F_}=\"TP\")+({Hh}=\"TP\"))>0))*({J}<>\"TP\"))","None expected"),
("Missed by all three",f"=SUMPRODUCT(({R}=\"Vulnerable\")*({F_}<>\"TP\")*({Hh}<>\"TP\")*({J}<>\"TP\")*({J}<>\"PARTIAL\"))",""),
]
for i,(a,f,d) in enumerate(extra,16):
    put(sm,i,[a,f,None,None,None,d])
sm["A21"]="Notes: counts are descriptive for 12 controlled fixtures. Alert counts are not quality measures; scanner rules for one condition are normalised to one detection. AI independent classification uses 3 identical runs per case."
sm["A21"].font=Font(name=F,italic=True); sm["A21"].alignment=Alignment(wrap_text=True); sm.merge_cells("A21:F22")

# ---------- Independent runs
ir=wb.create_sheet("Independent AI Runs")
cols=["Test case","PR","Run ID","Event","Started (UTC)","Commit","Condition (metadata)","Scanner findings supplied to AI","AI output status","Confirmed findings","Other findings (not in PR comment)","Operational score","Verdict"]
header(ir,1,cols,[8,6,14,16,20,9,13,13,14,55,55,9,22])
for i,(tc,md,ai) in enumerate(runs,2):
    if ai is None:
        conf,oth,score,verd,st="—","Output rejected by validator: FALSE_POSITIVE/INFORMATIONAL 'No azurerm_monitor_diagnostic_setting configured for the NSG' with a policy mapping to a non-confirmed finding","—","—","AI step failed"
    else:
        c=[f for f in ai["findings"] if f["status"]=="CONFIRMED"]; o=[f for f in ai["findings"] if f["status"]!="CONFIRMED"]
        conf="\n".join(f"{f['severity']} · {f['category']} · {f['title']}" for f in c) or "none"
        oth="\n".join(f"{f['status']} · {f['severity']} · {f['title']}" for f in o) or "none"
        score,verd,st=ai["operational_score"],ai["verdict"],"parsed"
    put(ir,i,[tc,int(md["pull_request_number"]),int(md["workflow_run_id"]),md["event_name"],md["started_at_utc"][:19].replace("T"," "),md["experiment_commit_sha"][:7],md["experiment_condition"],str(md["scanner_findings_supplied_to_ai"]).upper(),st,conf,oth,score,verd])
ir.freeze_panes="B2"; ir.auto_filter.ref=f"A1:M{len(runs)+1}"
k=len(runs)+3
ir.cell(row=k,column=1,value="Source: experiment-metadata.json and ai_output.json in each prom06-INDEPENDENT-<run_id> artifact (downloaded by AB, 13 Oct 2026). Model: claude-sonnet-4-6 (hard-coded in ai/review_runner.py; not recorded in metadata); no temperature set.").font=Font(name=F,italic=True)

# ---------- Scanner results
sr=wb.create_sheet("Scanner Results")
cols=["Test case","Checkov passed","Checkov failed","Checkov failed check IDs","tfsec findings (rule · severity · resource)","Identical across independent runs","Matches local reproduction","Matches earlier Condition C PR comment"]
header(sr,1,cols,[8,10,10,50,60,14,14,16])
for i in range(1,13):
    tc=f"TC{i:02d}"; d=f"{LOC}/{tc}"
    c=json.load(open(d+"/checkov.json")); c=c[0] if isinstance(c,list) else c
    t=json.load(open(d+"/tfsec.json")).get("results") or []
    put(sr,i+1,[tc,c["summary"]["passed"],c["summary"]["failed"],", ".join(x["check_id"] for x in c["results"]["failed_checks"]) or "none",
        "\n".join(f"{x['rule_id']} · {x['severity']} · {x['resource']}" for x in t) or "none","Yes (3/3; TC08 4/4)","Yes","Yes"])
sr.cell(row=15,column=1,value="Local reproduction: Checkov 3.3.1 (--framework terraform --output json) and tfsec v1.28.14 (--format json) on each indep-tc branch, 13 Oct 2026. Checkov OSS JSON reports no severity. Probe: changing TC02 source '0.0.0.0/0' to '*' makes CKV_AZURE_10 fail.").font=Font(name=F,italic=True)

# ---------- Hybrid runs
hr=wb.create_sheet("Hybrid Runs (Condition C)")
cols=["Test case","PR","Head commit","Run ID","Run attempts","AI result (PR comment)","Scanner IDs cited by AI","Notes"]
header(hr,1,cols,[8,6,10,14,9,55,30,50])
H_=[("TC01","#46","0003ccf","37804553671",1,"55/100 DO_NOT_MERGE: MEDIUM TAG-001, MEDIUM LOG-001 confirmed; context note on missing explicit NSG rules","none (no scanner context in code at that time)","Older fixture (no tags/logging); later baseline changed"),
("TC01","#55","aea09b1","37941043133",1,"Screenshot shows 100/100 clean, but PR changed only PROM06-TC-01.md so AI should have been skipped; comment posted 11 min before scanner comments","—","Not traceable to its own run — UNVERIFIED"),
("TC02","#56","1112b8b","37949921520",1,"70/100: CRITICAL SSH","AVD-AZU-0047, AVD-AZU-0050",""),
("TC03","#57","4d213c6","38240217182",1,"70/100: CRITICAL RDP","AVD-AZU-0048, AVD-AZU-0047",""),
("TC04","#58","4d2f880","38241873851",1,"70/100: HIGH HTTP","AVD-AZU-0047","PR closed unmerged"),
("TC05","#59","dad9df3","38243286069",1,"85/100: MEDIUM NET-001","CKV2_AZURE_31",""),
("TC06","#61","1ea2c15","38245294170",1,"85/100: MEDIUM NET-002","CKV_AZURE_183, CKV_AZURE_182","#60 was an earlier TC06 variant"),
("TC07","#63","de8549e","38281784140",1,"25/100: CRITICAL TLS, MEDIUM TAG-001, HIGH public network","AVD-AZU-0011, CKV_AZURE_44, CKV_AZURE_59, CKV2_AZURE_47/33/40","#62 web-app variant not used"),
("TC08","#64","7cb9014","38283490262",1,"85/100: MEDIUM LOG-001","none (scanners clean)",""),
("TC09","#65","eca6e99","38285110293",3,"85/100: MEDIUM hard-coded secret","none (scanners clean)","3 attempts; claimed value would be stored in state"),
("TC10","#66","dd05ba4","38287543507",2,"Attempt 1: 40/100, 2 × CRITICAL; attempt 2 (13 Oct): 70/100, 1 × HIGH","CKV_AZURE_39","Explains matrix severity conflict"),
("TC11","#67","960dac3","38288761497",2,"Both attempts 70/100: 1 × HIGH","CKV_AZURE_39 (listed as a 'policy mapping' in attempt 2)","Confluence heading 'AI-only (B)' incorrect"),
("TC12","#68","709316c","38289955785",1,"85/100: MEDIUM SEC-001 scope concern","AVD-AZU-0047","")]
for i,r in enumerate(H_,2): put(hr,i,r)
hr.cell(row=len(H_)+3,column=1,value="Runs on 9–12 Oct used workflow default Condition C with build_scanner_context() (added 9 Oct 12:47, commit 732083d), so the AI prompt contained normalised scanner findings. Run IDs/attempts from GitHub Actions API.").font=Font(name=F,italic=True)

# ---------- Decisions
dc=wb.create_sheet("Decisions & Caveats")
header(dc,1,["#","Item","Proposed treatment","AB decision"],[4,60,70,25])
D=[("TC12 tfsec AVD-AZU-0047 and AI SEC-001: count the shared-NSG / wildcard-destination scope as a separate weakness?","Keep as CONTEXT; report both readings (FP against TC12-GT-01 vs valid scope concern)."),
("TC08 AI result","Report PARTIAL: raised 4/4, confirmed 1/3 successful runs; treat as classification instability."),
("TC01 hybrid result","Report as UNVERIFIED; use independent TC01 (clean 3/3) as the baseline result; disclose baseline evolution."),
("TC11 construct","Count detections, but state the case does not isolate multi-file reasoning (wildcard alone triggers CKV_AZURE_39)."),
("AI TAG-001 findings on non-taggable resources (TC10, TC12)","Count as AI false positives (unrelated to target)."),
("Input asymmetry: only the AI receives ai/policy.md (NET-001, NET-002, LOG-001, SEC-001, TAG-001)","Report as a fairness/validity limitation, especially for TC05, TC06, TC08, TC12."),
("Severity comparison","Do not compare severities across tools; report AI severity variability only."),
("Ground truth source","Use Confluence v72 per-case ground truth (more specific than GT Matrix v1.0); note GT v1.0 TC04 wording differs."),
("Checkov TC02–TC04 misses","Report as reproducible rule-matching limitation with the local probe as evidence."),
("Close PRs #56–#81 without merging after sign-off","—")]
for i,(a,b) in enumerate(D,2):
    put(dc,i,[i-1,a,b,""]); dc.cell(row=i,column=4).fill=PatternFill("solid",fgColor="FFFF00")

# ---------- Draft copy
src=load_workbook("/root/.claude/uploads/8123617a-09e1-5f82-8592-b520c5c856e4/55793009-PROM06_TC01-TC12_Results_Matrix_Draft.xlsx").active
dm=wb.create_sheet("Draft Matrix (original)")
for r in src.iter_rows(values_only=True):
    dm.append(list(r))
for row in dm.iter_rows():
    for c in row: c.font=B; c.alignment=WR; c.border=BOR
for j in range(1,6): dm.cell(row=1,column=j).font=H; dm.cell(row=1,column=j).fill=HF; dm.column_dimensions[get_column_letter(j)].width=30
dm.cell(row=15,column=1,value="Unchanged copy of AB's draft (Results_Matrix_Draft.xlsx) for traceability. Superseded by 'Verified Matrix'.").font=Font(name=F,italic=True)
wb.save(OUT); print(OUT)
