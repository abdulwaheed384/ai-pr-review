"""Dissertation figures from the recalculated PROM06 workbook (values read with data_only=True)."""
import os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Patch
from openpyxl import load_workbook

P = "/home/claude/prom06-outputs/PROM06_TC01-TC12_Verified_Results_Matrix.xlsx"
OUT = "/home/claude/prom06-outputs/figures"
os.makedirs(OUT, exist_ok=True)
wb = load_workbook(P, data_only=True)

SURF = "#fcfcfb"; INK = "#0b0b0b"; INK2 = "#52514e"; GRID = "#e4e3df"
S1 = "#2a78d6"; S2 = "#eb6834"
GOOD, WARN, SER, CRIT, NEUT = "#0ca30c", "#fab219", "#ec835a", "#d03b3b", "#f0efec"
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "axes.edgecolor": INK2, "axes.labelcolor": INK2,
                     "xtick.color": INK2, "ytick.color": INK2, "figure.facecolor": SURF, "axes.facecolor": SURF,
                     "savefig.facecolor": SURF})


def clean(ax, grid_axis="x"):
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.grid(axis=grid_axis, color=GRID, linewidth=0.8); ax.set_axisbelow(True)


def save(fig, name):
    fig.savefig(f"{OUT}/{name}.png", dpi=300, bbox_inches="tight"); plt.close(fig)


# ---------- Figure: detection outcome matrix ----------
vm = wb["Verified Matrix"]; met = wb["Metrics"]; ro = wb["AI Run Outcomes"]
names = {"TC01": "Secure baseline", "TC02": "Public SSH", "TC03": "Public RDP", "TC04": "Public HTTP", "TC05": "NSG not associated",
         "TC06": "Public DNS resolver", "TC07": "Weak TLS (storage)", "TC08": "Missing NSG diagnostics", "TC09": "Secret in local",
         "TC10": "Wildcard custom role", "TC11": "Role vs documented need", "TC12": "Legitimate public HTTPS",
         "TC13": "Contributor vs Reader need", "TC14": "SSH from one address", "TC15": "Password in variable default"}
runs = {}
for r in ro.iter_rows(min_row=5, values_only=True):
    if r[0]:
        runs.setdefault(r[0], []).append(r)
rows = []
for i in range(15):
    tc = met.cell(row=5 + i, column=1).value
    ck, tf, ai = met.cell(row=5 + i, column=3).value, met.cell(row=5 + i, column=4).value, met.cell(row=5 + i, column=6).value
    ct = met.cell(row=5 + i, column=2).value
    rr = runs[tc]
    if ct == "Vulnerable":
        k = sum(1 for x in rr if x[9] == "Y"); ai_txt = f"{ai}\n{k}/{len(rr)} runs"
    else:
        k = sum(1 for x in rr if (x[4] or 0) == 0); ai_txt = f"{ai}\n{k}/{len(rr)} runs clean"
    hy = vm.cell(row=2 + i, column=12).value if i < 12 else "not run"
    rows.append((tc, names[tc], ct, ck, tf, ai_txt, ai, hy))

COL = {"TP": GOOD, "TN": GOOD, "FN": CRIT, "PARTIAL": WARN, "CONTEXT": SER, "UNVERIFIED": NEUT, "not run": NEUT}
fig, ax = plt.subplots(figsize=(9.6, 8.6))
cols = ["Checkov\n3.3.1", "tfsec\n1.28.14", "AI\nindependent", "AI hybrid\n(Condition C)"]
for i, (tc, nm, ct, ck, tf, ai_txt, ai, hy) in enumerate(rows):
    y = len(rows) - 1 - i
    for j, (cls, txt) in enumerate(((ck, ck), (tf, tf), (ai, ai_txt), (hy, hy))):
        ax.add_patch(plt.Rectangle((j + 0.03, y + 0.06), 0.94, 0.88, color=COL[cls], linewidth=0))
        ax.text(j + 0.5, y + 0.5, txt, ha="center", va="center", fontsize=7.8, linespacing=1.1,
                color="white" if cls == "FN" else INK, fontweight="bold" if cls in ("TP", "FN") else "normal")
    lab = f"{tc}  {nm}"
    ax.text(-0.08, y + 0.5, lab, ha="right", va="center", fontsize=9, color=INK)
    ax.text(4.08, y + 0.5, ct, ha="left", va="center", fontsize=8, color=INK2)
for j, c in enumerate(cols):
    ax.text(j + 0.5, len(rows) + 0.15, c, ha="center", va="bottom", fontsize=8.5, color=INK, fontweight="bold")
ax.text(4.08, len(rows) + 0.15, "Ground\ntruth", ha="left", va="bottom", fontsize=8.5, color=INK, fontweight="bold")
ax.set_xlim(-2.6, 5.0); ax.set_ylim(-2.9, len(rows) + 1.1); ax.axis("off")
leg = [Patch(color=GOOD, label="Correct (TP / TN)"), Patch(color=CRIT, label="Missed (FN)"),
       Patch(color=WARN, label="Partial (raised, not consistently confirmed)"), Patch(color=SER, label="Context-dependent"),
       Patch(color=NEUT, label="Unverified / not run")]
ax.legend(handles=leg, loc="upper center", bbox_to_anchor=(0.5, 0.115), ncol=3, frameon=False, fontsize=8.5)
ax.text(-2.5, -2.85, "AI independent cells: case class, then runs confirming the target\n"
        "(vulnerable) or runs with no confirmed finding (non-vulnerable). Scanner rules normalised to one detection per case.", fontsize=7.5, color=INK2, va="bottom")
save(fig, "fig_detection_matrix")

# ---------- Figure: detection rate by tool ----------
def blk(top):
    return [met.cell(row=top, column=c).value for c in (3, 4, 5, 6)]
# locate strict-rate rows by label
rate_rows = [r for r in range(20, 60) if met.cell(row=r, column=1).value == "Detection rate, strict (TP ÷ vulnerable)"]
a, b = blk(rate_rows[0]), blk(rate_rows[1])
tools = ["Checkov", "tfsec", "Checkov ∪ tfsec", "AI independent"]
fig, ax = plt.subplots(figsize=(7.2, 3.6))
import numpy as np
x = np.arange(len(tools)); w = 0.36
b1 = ax.bar(x - w / 2 - 0.01, a, w, color=S1, label="TC01–TC12, verified (10 vulnerable cases)")
b2 = ax.bar(x + w / 2 + 0.01, b, w, color=S2, label="TC01–TC15 (12 vulnerable cases)")
for bars in (b1, b2):
    for r_ in bars:
        ax.text(r_.get_x() + r_.get_width() / 2, r_.get_height() + 0.02, f"{r_.get_height():.0%}", ha="center", va="bottom", fontsize=8.5, color=INK)
ax.set_xticks(x, tools); ax.set_ylim(0, 1.1); ax.yaxis.set_major_formatter(matplotlib.ticker.PercentFormatter(1.0))
ax.set_ylabel("Vulnerable cases detected (strict)"); clean(ax, "y")
ax.legend(frameon=False, fontsize=8.5, loc="upper left")
ax.text(0, -0.24, "Strict: AI PARTIAL (TC08) counted as a miss; lenient AI rate is 100% in both sets. Descriptive only: constructed cases, not a sample.",
        transform=ax.transAxes, fontsize=7.5, color=INK2)
save(fig, "fig_detection_rates")

# ---------- Figure: AI quality by dimension ----------
q = wb["AI Quality Scoring"]
mrow = next(r for r in range(30, 60) if (q.cell(row=r, column=1).value or "").startswith("Mean — vulnerable"))
dims = ["Detection", "Explanation accuracy", "Impact", "Remediation", "Context awareness", "Developer usefulness"]
vals = [q.cell(row=mrow, column=c).value for c in range(7, 13)]
fig, ax = plt.subplots(figsize=(6.8, 3.2))
yy = np.arange(len(dims))[::-1]
ax.barh(yy, vals, height=0.55, color=S1)
for y_, v in zip(yy, vals):
    ax.text(v + 0.03, y_, f"{v:.2f}", va="center", fontsize=8.5, color=INK)
ax.set_yticks(yy, dims); ax.set_xlim(0, 2.2); ax.set_xticks([0, 0.5, 1, 1.5, 2])
ax.set_xlabel("Mean score (0–2), 37 independent runs on vulnerable cases (TC02–TC11, TC13, TC15)"); clean(ax, "x")
ax.text(0, -0.36, "Scored by the AI research assistant against predefined operational rules; inter-rater check by the researcher pending.",
        transform=ax.transAxes, fontsize=7.5, color=INK2)
save(fig, "fig_ai_quality")

# ---------- Figure: timing ----------
tm = wb["Timing"]
med = next(r for r in range(40, 70) if tm.cell(row=r, column=1).value == "Median")
lab_t = [("Checkov install", "F"), ("Checkov scan", "G"), ("tfsec scan", "H"), ("AI review (incl. API call)", "I"), ("Whole job", "E"), ("PR opened → comments posted", "L")]
vals = [tm[f"{c}{med}"].value for _, c in lab_t]
fig, ax = plt.subplots(figsize=(6.8, 3.0))
yy = np.arange(len(lab_t))[::-1]
ax.barh(yy, vals, height=0.55, color=S1)
for y_, v in zip(yy, vals):
    ax.text(v + 1, y_, f"{v:.1f} s", va="center", fontsize=8.5, color=INK)
ax.set_yticks(yy, [l for l, _ in lab_t]); ax.set_xlim(0, max(vals) * 1.18)
ax.set_xlabel("Median seconds (46 independent runs; PR latency from 15 pull_request runs)"); clean(ax, "x")
save(fig, "fig_timing")

# ---------- Figure: remediation validation ----------
rv = wb["Remediation Validation"]
srow = next(r for r in range(15, 40) if (rv.cell(row=r, column=1).value or "").startswith("Cases with an AI remediation"))
get = lambda k: rv.cell(row=srow + k, column=6).value
n = get(0)
items = [("Target resolved vs ground truth", get(1)), ("Passed terraform fmt + validate", get(2)),
         ("Confirmed clean by the scanner that had flagged it", get(3)), ("AI re-review accepted its own fix", get(5))]
fig, ax = plt.subplots(figsize=(6.8, 2.6))
yy = np.arange(len(items))[::-1]
ax.barh(yy, [n] * len(items), height=0.55, color=NEUT)
ax.barh(yy, [v for _, v in items], height=0.55, color=S1)
for y_, (_, v) in zip(yy, items):
    ax.text(v + 0.15, y_, f"{v} of {n}", va="center", fontsize=8.5, color=INK)
ax.set_yticks(yy, [l for l, _ in items]); ax.set_xlim(0, n + 2.5); ax.set_xlabel(f"AI-recommended fixes applied (n = {n}; one run each)")
clean(ax, "x")
fig.text(0.01, -0.08, f"Scanner confirmation impossible for {get(4)} cases the scanners never detected. AI rejected 2 fixes (TC02, TC03) while stale\n"
        "'deliberately unrestricted' comments remained; it accepted both once only the comment was corrected. TC10 deletion not re-reviewed by the AI.",
        fontsize=7.5, color=INK2)
save(fig, "fig_remediation")
print(sorted(os.listdir(OUT)))
