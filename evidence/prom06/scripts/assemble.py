"""Assemble the full PROM06 dissertation markdown (front matter, chapters, references, appendices)."""
import re, json, glob, subprocess, os
from openpyxl import load_workbook

CH = "/home/claude/prom06-outputs/chapters"
OUT = "/home/claude/prom06-outputs/final"
os.makedirs(OUT, exist_ok=True)
REPO = "/home/claude/ai-pr-review"
WB = "/home/claude/prom06-outputs/PROM06_TC01-TC12_Verified_Results_Matrix.xlsx"
PAGEBREAK = '\n```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```\n'

def split_refs(md):
    if "## References cited in this chapter" in md:
        body, refs = md.split("## References cited in this chapter")
        entries = [e.strip() for e in refs.strip().split("\n\n") if e.strip()]
    else:
        body, entries = md, []
    return body.rstrip() + "\n", entries

def words(md):
    body = [l for l in md.splitlines() if l.strip() and not l.startswith('|') and not l.startswith('Table ')
            and not l.startswith('![') and not l.startswith('# ') and not l.startswith('```') and not l.startswith('<w:')]
    return len(re.sub(r'[#*`]', '', ' '.join(body)).split())

# ---------------------------------------------------------------- chapters
chapters, allrefs = [], []
for f in ["ch1", "ch2", "ch3", "ch4", "ch5", "ch6"]:
    b, r = split_refs(open(f"{CH}/{f}.md").read())
    b = b.replace("](../figures/", "](/home/claude/prom06-outputs/figures/")
    chapters.append(b); allrefs += r
abstract_md = open(f"{CH}/abstract.md").read().replace("# Abstract", "").strip()
abstract_para = " ".join(p.strip().replace("\n", " ") for p in abstract_md.split("\n\n") if p.strip())
wc = words("Abstract\n" + abstract_para) + sum(words(c) for c in chapters)

# ---------------------------------------------------------------- references (dedupe)
def key(e):
    m = re.match(r"(.+?)\((\d{4}[a-z]?|n\.d\.)\)\s*'?\*?([^'*.]{0,40})", e)
    return (m.group(1).split(',')[0].strip().lower(), m.group(2), m.group(3).strip().lower()[:25]) if m else (e[:40],)
PREFER_SHORT = ("johnson",)  # uncertain DOI variants — keep the version without DOI
refs = {}
for e in allrefs:
    k = key(e)
    if k not in refs: refs[k] = e
    else:
        cur = refs[k]
        if k[0] in PREFER_SHORT: refs[k] = min(cur, e, key=len)
        else: refs[k] = max(cur, e, key=len)
ref_list = sorted(refs.values(), key=lambda e: re.sub(r'[^a-z0-9 ]', '', e.lower()))

# ---------------------------------------------------------------- appendix A: proposal text
src = open("/home/claude/prom06-sources/06_PROM04_Proposal_and_Slides.md").read().splitlines()
start = next(i for i, l in enumerate(src) if l.strip() == "1 Introduction" and i > 30)
end = next(i for i, l in enumerate(src) if l.startswith("--- Slide 1"))
lines = [l.strip() for l in src[start:end] if l.strip() and not l.startswith("#") and "PROM04 ASSIGNMENT-02" not in l and l.strip() != "AI-ASSISTED CODE REVIEW IN CI/CD PIPELINE"]
paras, cur = [], ""
head_re = re.compile(r"^(\d(\.\d)?)\s+[A-Z][A-Za-z ,&/\-]+$")
for l in lines:
    if head_re.match(l) and len(l) < 80:
        if cur: paras.append(cur); cur = ""
        lvl = "####" if "." in l.split()[0] else "###"
        paras.append(f"{lvl} A.{l}"); continue
    cur = (cur + " " + l).strip()
    if re.search(r"[.:;)]$", l) and not l.endswith("e.g.") and not l.endswith("et al."):
        paras.append(cur); cur = ""
if cur: paras.append(cur)
proposal = "\n\n".join(paras).replace("### A.7 References", "### A.7 References (as submitted in PROM04)")
proposal = re.sub(r"### A\.6 Social, Ethical, Professional and Legal\n\nConsiderations ?", "### A.6 Social, Ethical, Professional and Legal Considerations\n\n", proposal)

# ---------------------------------------------------------------- data for appendices
wb = load_workbook(WB, data_only=True)
runs_tsv = [l.split("\t") for l in open(f"{REPO}/../prom06-independent-evidence/independent-runs.tsv").read().splitlines()[1:]] \
    if os.path.exists(f"{REPO}/../prom06-independent-evidence/independent-runs.tsv") else []
runs_tsv = [l.split("\t") for l in open("/home/claude/prom06-independent-evidence/independent-runs.tsv").read().splitlines()[1:]]
for l in open("/home/claude/prom06-step3/tc13-15-runs.tsv").read().splitlines():
    p = l.split("\t"); runs_tsv.append([p[1], p[2], p[0], p[3], p[4], p[5], p[6]])
by_tc = {}
for br, ev, rid, sha, att, ts, concl in runs_tsv:
    tc = br.replace("indep-tc", "TC"); by_tc.setdefault(tc, {"sha": sha, "runs": []})["runs"].append((rid, ev, concl))
names = {"TC01": "Secure baseline", "TC02": "Unrestricted SSH", "TC03": "Unrestricted RDP", "TC04": "Public HTTP", "TC05": "Missing NSG association",
         "TC06": "Public DNS resolver", "TC07": "Weak TLS (storage)", "TC08": "Missing NSG diagnostics", "TC09": "Secret in local",
         "TC10": "Wildcard custom role", "TC11": "Role vs documented need", "TC12": "Legitimate public HTTPS", "TC13": "Contributor vs Reader need",
         "TC14": "SSH from one address", "TC15": "Password in variable default"}

def md_table(header, rows, widths=None):
    widths = widths or [8] * len(header)
    out = "| " + " | ".join(header) + " |\n|" + "|".join("-" * w for w in widths) + "|\n"
    for r in rows: out += "| " + " | ".join("" if c is None else str(c).replace("|", "/").replace("\n", " ") for c in r) + " |\n"
    return out

# D.1 evidence index
d1 = []
for i, tc in enumerate(sorted(by_tc)):
    pr = 70 + i
    d1.append([tc, names[tc], f"indep-{tc.lower()}", by_tc[tc]["sha"], f"#{pr}", ", ".join(r[0] + (" (failed)" if r[2] == "failure" else "") for r in by_tc[tc]["runs"])])
rem = []
prmap = {"rem-tc02": 85, "rem-tc03": 86, "rem-tc04": 87, "rem-tc05": 88, "rem-tc06": 89, "rem-tc07": 90, "rem-tc08": 91, "rem-tc09": 92, "rem-tc10": 93,
         "rem-tc10tags": 94, "rem-tc11": 95, "rem-tc13": 96, "rem-tc15": 97, "rem-tc02b": 98, "rem-tc03b": 99, "rem-tc04b": 100}
for f in ("/home/claude/prom06-step3/rem-runs.tsv", "/home/claude/prom06-step3/rem-runs-b.tsv"):
    for l in open(f).read().splitlines():
        rid, br, sha, concl, ts = l.split("\t"); rem.append([br, sha, f"#{prmap[br]}", rid, concl])
rem.sort(key=lambda r: int(r[2][1:]))

# E tables from workbook
m = wb["Metrics"]
e1 = []
vm = wb["Verified Matrix"]
for i in range(15):
    tc = m.cell(row=5 + i, column=1).value
    hy = vm.cell(row=2 + i, column=12).value if i < 12 else "not run"
    e1.append([tc, names[tc], m.cell(row=5 + i, column=2).value, m.cell(row=5 + i, column=3).value, m.cell(row=5 + i, column=4).value,
               m.cell(row=5 + i, column=6).value, hy])
q = wb["AI Quality Scoring"]
ro = wb["AI Run Outcomes"]
outcomes = {str(r[2]): r for r in ro.iter_rows(min_row=5, values_only=True) if r[0]}
e2 = []
for r in q.iter_rows(min_row=2, values_only=True):
    if not r[0] or not str(r[0]).startswith("TC") or r[1] is None: continue
    o = outcomes.get(str(r[2]))
    sc = lambda v: "–" if v in (None, "n/a") else v
    e2.append([r[0], r[1], r[2], o[9] if o else "", o[4] if o else "", o[8] if o else "", sc(r[6]), sc(r[7]), sc(r[8]), sc(r[9]), sc(r[10]), sc(r[11]),
               f"{round(r[14]*100)}%" if isinstance(r[14], (int, float)) else ""])
cf = wb["AI Confirmed Findings"]
e3 = [[r[0], r[1], r[4], r[5], r[6], r[7]] for r in cf.iter_rows(min_row=6, values_only=True) if r[0] and r[6] in ("FALSE", "DUPLICATE", "EXTRA-VALID")]
sr = wb["Scanner Results"]
e4 = [[r[0], r[1], r[2], r[3], r[4]] for r in sr.iter_rows(min_row=2, max_row=13, values_only=True)]
loc = "/home/claude/prom06-local-scans"
for tc in ("TC13", "TC14", "TC15"):
    c = json.load(open(f"{loc}/{tc}/checkov.json")); c = c[0] if isinstance(c, list) else c
    t = json.load(open(f"{loc}/{tc}/tfsec.json")).get("results") or []
    e4.append([tc, c["summary"]["passed"], c["summary"]["failed"], "none", "none" if not t else ", ".join(x["long_id"] for x in t)])
dc = wb["Decisions & Caveats"]
e5 = [[r[0], r[1], r[2], r[3]] for r in dc.iter_rows(min_row=2, values_only=True) if r[0] is not None and r[1]]
tm = wb["Timing"]
med = next(r for r in range(40, 70) if tm.cell(row=r, column=1).value == "Median")
e7 = []
for lab, rr in (("Median", med), ("Minimum", med + 1), ("Maximum", med + 2)):
    e7.append([lab] + [tm[f"{c}{rr}"].value for c in "EFGHIL"])

# ---------------------------------------------------------------- appendix C code extracts
def code(path, lang=""):
    return f"```{lang}\n{open(path).read().rstrip()}\n```\n"
runner = open(f"{REPO}/ai/review_runner.py").read().splitlines()
ctx_fn = "\n".join(runner[27:64])

# ---------------------------------------------------------------- appendix B ground truth TC13-15
gt1315 = subprocess.run(["git", "-C", REPO, "show", "prom06-evidence:evidence/prom06/ground-truth/TC13-TC15_Ground_Truth.md"], capture_output=True, text=True).stdout
gt1315 = re.sub(r"^# .*\n", "", gt1315).replace("\n## ", "\n### ")

# ---------------------------------------------------------------- appendix F screenshots
SS = "/home/claude/prom06-step3/screenshots-src"
FIMG = f"{OUT}/appendix_f"; os.makedirs(FIMG, exist_ok=True)
caps = [("PR71", [2, 3], "S1. TC02 pull request #71: scanner and AI comments"),
        ("PR78", [2, 3], "S2. TC09 pull request #78: scanners clean, AI reports the embedded secret"),
        ("S3_run38404605125", [1, 2], "S3. TC08 run 38404605125 (pull request #77): AI step fails validation"),
        ("PR81", [2, 3], "S4. TC12 pull request #81: tfsec alert and AI block on the intended HTTPS rule"),
        ("PR82", [2, 3], "S5a. TC13 pull request #82: scanners clean, AI reports Contributor at subscription scope"),
        ("PR84", [2, 3], "S5b. TC15 pull request #84: scanners clean, AI reports the hard-coded password"),
        ("PR85", [3], "S6a. Remediation pull request #85 (TC02 fix, stale comment)"),
        ("PR98", [3], "S6b. Remediation pull request #98 (TC02 fix, comment corrected)"),
        ("S7_run38412955313", [1, 2], "S7. Run 38412955313 (pull request #94): AI tag recommendation fails terraform validate"),
        ("S8_run38404584887", [1], "S8. TC02 run 38404584887: workflow job steps and durations")]
fig_md = ""
for base, pages, cap in caps:
    pdf = f"{SS}/{base}.pdf"
    for p in pages:
        stem = f"{FIMG}/{base}_p{p}"
        subprocess.run(["pdftoppm", "-jpeg", "-jpegopt", "quality=80", "-r", "90", "-f", str(p), "-l", str(p), "-singlefile", pdf, stem], check=True)
        fig_md += f"*{cap} (page {p})*\n\n![]({stem}.jpg){{width=15cm}}\n\n"

# ---------------------------------------------------------------- assemble
front = f"""---
title: "Evaluating AI-Driven Security Review in Modern DevSecOps Pipelines"
---

```{{=openxml}}
<w:p><w:pPr><w:spacing w:before="1800"/></w:pPr></w:p>
```

::: {{custom-style="TitleBlock"}}
Evaluating AI-Driven Security Review in Modern DevSecOps Pipelines

An empirical comparison of an AI reviewer with Checkov and tfsec for Terraform pull requests

Abdul Waheed

Student number: 209333039

Supervisor: [Supervisor name]

A dissertation submitted in partial fulfilment of the requirements for the degree of MSc Cybersecurity

University of Sunderland

[Faculty name], School of Computer Science

October 2026
:::
{PAGEBREAK}
# Declaration {{.unnumbered}}

I declare the following:

(1) that the material contained in this dissertation is the end result of my own work and that due acknowledgement has been given in the bibliography and references to ALL sources be they printed, electronic or personal.

(2) the Word Count of this Dissertation is {wc:,}

(3) that unless this dissertation has been confirmed as confidential, I agree to an entire electronic copy or sections of the dissertation to being placed on the eLearning Portal, if deemed appropriate, to allow future students the opportunity to see examples of past dissertations. I understand that if displayed on the eLearning Portal it would be made available for no longer than five years and that students would be able to print off copies or download.

(4) I agree to my dissertation being submitted to a plagiarism detection service, where it will be stored in a database and compared against work submitted from this or any other Department or from other institutions using the service.

In the event of the service detecting a high degree of similarity between content within the service this will be reported back to my supervisor and second marker, who may decide to undertake further investigation that may ultimately lead to disciplinary actions, should instances of plagiarism be detected.

(5) I have read the University of Sunderland Policy Statement on Ethics in Research and I confirm that ethical issues have been considered, evaluated and appropriately addressed in this research.

SIGNED: ............................................

DATE: ..............................................
{PAGEBREAK}
# Abstract {{.unnumbered}}

{abstract_para}
{PAGEBREAK}
# Contents {{.unnumbered}}

```{{=openxml}}
<w:p><w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r><w:r><w:instrText xml:space="preserve"> TOC \\o "1-2" \\h \\z \\u </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>Table of contents</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>
```
{PAGEBREAK}
# List of Figures {{.unnumbered}}

```{{=openxml}}
<w:p><w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r><w:r><w:instrText xml:space="preserve"> TOC \\h \\z \\t "Image Caption,1" </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>List of figures</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>
```

# List of Tables {{.unnumbered}}

```{{=openxml}}
<w:p><w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r><w:r><w:instrText xml:space="preserve"> TOC \\h \\z \\t "Table Caption,1" </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>List of tables</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>
```
"""

body = ""
for c in chapters:
    body += PAGEBREAK + c + "\n"

refs_md = PAGEBREAK + "# 7 Reference List\n\n" + "\n\n".join(ref_list) + "\n"

app = PAGEBREAK + "# Appendix A. Research Proposal {.unnumbered}\n\n"
app += "*Text of the PROM04 research proposal, 'AI-Assisted Code Review in CI/CD Pipeline', submitted 20 April 2026, reproduced as submitted (text only; the proposal's own cover page and table of contents are omitted). The research focus was refined in PROM05 and PROM06 Assignment 1 from a comparison with manual review to a comparison with IaC static analysis tools.*\n\n"
app += proposal + "\n"

app += PAGEBREAK + "# Appendix B. Experimental Plan and Ground-Truth Matrix {.unnumbered}\n\n"
app += "## B.1 Experimental conditions and ground-truth rules\n\n"
app += ("The experimental master plan (version 1.0) defined three conditions: A, traditional scanning (Checkov and tfsec); B, AI only; and C, hybrid. "
        "The ground-truth matrix (TC01–TC12, version 1.0, revised in the test-case evidence page, version 72) set the following rules: ground truth is established before tool outputs are reviewed; each case is classified as vulnerable, secure or intentionally acceptable independently of the AI; uncertain benchmark mappings are recorded as 'to be verified'; tool expectations are hypotheses until verified; overlapping scanner rules are normalised; and the AI is judged against the documented ground truth, not against agreement with another tool.\n\n")
app += "## B.2 Test cases TC01–TC15\n\nTable B.1: Test-case ground truth and expected secure state\n\n"
secure = {"TC01": "No change required", "TC02": "Source restricted to an approved administration address or private path", "TC03": "Source restricted, or private management path (e.g. Bastion)",
          "TC04": "No unrestricted public HTTP; HTTPS where a public service is needed", "TC05": "NSG associated with the subnet", "TC06": "Azure-provided or internal DNS",
          "TC07": "min_tls_version TLS1_2", "TC08": "Diagnostic setting forwarding NetworkSecurityGroupEvent and NetworkSecurityGroupRuleCounter to Log Analytics",
          "TC09": "No secret literal in code; secret store or sensitive variable", "TC10": "Least-privilege actions at the narrowest scope", "TC11": "Only the two documented read actions",
          "TC12": "Intended exposure retained; no blocking finding", "TC13": "Reader at resource-group scope", "TC14": "No change required", "TC15": "No default; value supplied at runtime from a secret store; sensitive = true"}
gt = {"TC01": "Secure", "TC12": "Legitimate", "TC14": "Secure"}
app += md_table(["ID", "Case", "Ground truth", "Expected secure state"], [[tc, names[tc], gt.get(tc, "Vulnerable"), secure[tc]] for tc in sorted(names)], [9, 24, 12, 45])
app += "\n## B.3 Ground truth for TC13–TC15 (committed before the fixtures, commit cbda52c)\n\n" + gt1315 + "\n"

app += PAGEBREAK + "# Appendix C. Implementation and Configuration Extracts {.unnumbered}\n\n"
app += "Full source: repository `abdulwaheed384/ai-pr-review`, branch `main` at merge commit 3a79020 (pull request #69).\n\n"
app += "## C.1 Organisational policy supplied to the AI (`ai/policy.md`)\n\n" + code(f"{REPO}/ai/policy.md", "markdown")
app += "\n## C.2 Operational scoring configuration (`ai/scoring.json`)\n\n" + code(f"{REPO}/ai/scoring.json", "json")
app += "\n## C.3 Context assembly in the review runner (`ai/review_runner.py`, lines 28–64)\n\n```python\n" + ctx_fn + "\n```\n"
app += "\n## C.4 Review instructions (`ai/prompt.txt`)\n\n" + code(f"{REPO}/ai/prompt.txt")
app += "\n## C.5 Workflow (`.github/workflows/terraform.yml`)\n\n" + code(f"{REPO}/.github/workflows/terraform.yml", "yaml")

app += PAGEBREAK + "# Appendix D. Evidence Index for TC01–TC15 {.unnumbered}\n\n"
app += ("All raw evidence is preserved, unaltered, on branch `prom06-evidence` of the repository under `evidence/prom06/`: run artefacts (`independent-runs/<run_id>/`), "
        "run indexes, pull-request comments (`pr-comments/`), local scanner reproductions (`local-scans/`), step-3 data (`step3/`), ground truth, figures, page captures "
        "(`screenshots/`) and the scripts that produce every table and figure (`scripts/`).\n\n")
app += "Table D.1: Independent runs by test case\n\n" + md_table(["ID", "Case", "Branch", "Commit", "PR", "Run IDs"], d1, [9, 20, 13, 9, 6, 43])
app += "\nTable D.2: Remediation-validation branches and runs\n\n" + md_table(["Branch", "Commit", "PR", "Run ID", "Result"], rem, [14, 9, 5, 13, 8])
app += "\nEarlier scanner-informed (Condition C) runs used pull requests #46 and #55–#68; their comments are archived in `pr-comments/condition-c/`.\n"

app += PAGEBREAK + "# Appendix E. Results, Metrics and Scoring Sheets {.unnumbered}\n\n"
app += "Source: results workbook `PROM06_TC01-TC12_Verified_Results_Matrix.xlsx` (evidence branch), in which every summary figure is a spreadsheet formula over the run-level data.\n\n"
app += "Table E.1: Case-level classification\n\n" + md_table(["ID", "Case", "Type", "Checkov", "tfsec", "AI independent", "AI hybrid (C)"], e1, [9, 22, 12, 9, 9, 12, 12])
app += "\nTable E.2: Scanner results\n\n" + md_table(["ID", "Checkov passed", "Checkov failed", "Checkov failed checks", "tfsec findings"], e4, [9, 9, 9, 30, 40])
app += "\nTable E.3: AI runs — target confirmation, confirmed findings and rubric scores (D detection, E explanation, I impact, R remediation, C context, U usefulness; '–' not scored for non-vulnerable cases)\n\n"
app += md_table(["TC", "Run", "Run ID", "Target confirmed", "Confirmed findings", "False", "D", "E", "I", "R", "C", "U", "Score"], e2, [7, 5, 14, 9, 9, 7, 4, 4, 4, 4, 4, 4, 8])
app += "\nTable E.4: Confirmed AI findings other than the target\n\n" + md_table(["TC", "Run", "Severity", "Title", "Label", "Justification"], e3, [8, 5, 9, 30, 12, 36])
app += "\nTable E.5: Recorded decisions on evidence conflicts (agreed with the researcher, 13 October 2026)\n\n" + md_table(["#", "Issue", "Treatment", "Decision"], e5, [4, 30, 40, 14])
app += "\nTable E.6: Pipeline timing across 46 independent runs (seconds)\n\n" + md_table(["Statistic", "Whole job", "Install Checkov", "Run Checkov", "Run tfsec", "AI review", "PR → comments"], e7, [8, 7, 7, 7, 7, 7, 7])

app += PAGEBREAK + "# Appendix F. Pipeline Logs and Page Captures {.unnumbered}\n\n"
app += "## F.1 Validation failure of the AI's tag recommendation (run 38412955313)\n\n```\n" + open("/home/claude/prom06-step3/rem-logs/38412955313-terraform-validate.txt").read().strip() + "\n```\n"
app += "\n## F.2 TC08 run 1 (run 38404605125): AI output rejected by the runner\n\n```\nAI review failed: Policy mapping 1 must refer to a matching finding\nError: Process completed with exit code 1.\n```\n"
app += ("\n## F.3 GitHub page captures\n\nPages printed to PDF from GitHub on 13 October 2026 (18:22–18:29 BST); each page footer shows its URL and print time. "
        "Content was cross-checked against the comments and logs retrieved through the GitHub API. An artefact re-downloaded at 18:31 was byte-identical to the copy archived earlier the same day.\n\n")
app += fig_md

app += PAGEBREAK + "# Appendix G. Project Planning {.unnumbered}\n\n"
app += "Table G.1: PROM05 milestones compared with the project as executed\n\n"
g = [["M1 Research questions confirmed", "Week 2", "Refined in PROM06 Assignment 1 (submitted 23 July 2026)"],
     ["M2 Literature review complete", "Week 3", "Assignment 1 review; reworked as Chapter 2 in October 2026"],
     ["M3 Ethical approval", "Week 4", "Not required: no human participants or personal data"],
     ["M4 Pipeline operational", "Week 9", "Repository from 27 April 2026; experimental workflow from 8 October 2026; independent workflow from 13 October 2026 (PR #69)"],
     ["M5 Data collection complete", "Week 13", "Scanner-informed runs 9–12 October; independent reruns, TC13–TC15 and remediation validation 13 October 2026. Developer sessions and interviews not conducted"],
     ["M6 Analysis complete", "Week 14", "13 October 2026 (results workbook, evidence register)"],
     ["M7 Draft submission", "Week 15", "Full draft 13 October 2026"],
     ["M8 Final submission", "Week 16", "Planned 25 October 2026 (deadline 26 October 2026)"]]
app += md_table(["PROM05 milestone", "Planned", "Actual"], g, [26, 8, 50])
app += ("\nThe main departures from the PROM05 plan were the omission of developer sessions and interviews (and therefore of thematic analysis and ethics review), "
        "the use of static review without Azure deployment, the discovery and correction of scanner-informed AI runs, and the addition of TC13–TC15, repeated runs and remediation validation.\n")

full = front + body + refs_md + app
full = re.sub(r"(?m)^(Table [0-9A-G]+\.\d+: .+)$", r": \1", full)
open(f"{OUT}/PROM06_Dissertation_full.md", "w").write(full)
print("word count", wc, "refs", len(ref_list))
