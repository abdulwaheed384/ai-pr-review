from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
F="Arial"; H=Font(name=F,bold=True,color="FFFFFF"); HF=PatternFill("solid",fgColor="1F3864")
B=Font(name=F); BB=Font(name=F,bold=True); TI=Font(name=F,bold=True,size=14); IT=Font(name=F,italic=True); BLUE=Font(name=F,color="0000FF")
thin=Side(style="thin",color="BFBFBF"); BOR=Border(left=thin,right=thin,top=thin,bottom=thin); WR=Alignment(wrap_text=True,vertical="top")
wb=Workbook()
def table(ws,r0,cols,widths,rows,datecols=(),statuscol=None,inputcols=()):
    for j,(c,w) in enumerate(zip(cols,widths),1):
        x=ws.cell(row=r0,column=j,value=c); x.font=H; x.fill=HF; x.border=BOR; x.alignment=Alignment(wrap_text=True,vertical="center")
        ws.column_dimensions[chr(64+j)].width=w
    for i,r in enumerate(rows,r0+1):
        for j,v in enumerate(r,1):
            c=ws.cell(row=i,column=j,value=v); c.border=BOR; c.alignment=WR; c.font=BLUE if j in inputcols else B
            if j in datecols and isinstance(v,date): c.number_format="ddd dd mmm yyyy"
    last=r0+len(rows)
    if statuscol:
        L=chr(64+statuscol); rng=f"{L}{r0+1}:{L}{last}"
        for t,col in (("Done","C6EFCE"),("In progress","FFEB9C"),("Next","FFEB9C"),("To do","FFFFFF"),("Pending AB","FFC7CE"),("Blocked","FFC7CE")):
            ws.conditional_formatting.add(rng,CellIsRule(operator="equal",formula=[f'"{t}"'],fill=PatternFill("solid",fgColor=col)))
        dv=DataValidation(type="list",formula1='"Done,In progress,Next,To do,Pending AB,Blocked"',allow_blank=True); ws.add_data_validation(dv); dv.add(rng)
    return last
D=lambda d:date(2026,10,d)
# ---- Plan
ws=wb.active; ws.title="Plan"
ws["A1"]="PROM06 Assignment 2 — project plan"; ws["A1"].font=TI
ws["A2"]="Updated 13 Oct 2026, 16:00. Deadline assumed 26 Oct 2026 23:59 (Canvas export — confirm on live Canvas); target submission 25 Oct. Blue cells (dates, owner, status, notes) are editable; summary counts are formulas."
ws["A2"].font=IT; ws["A2"].alignment=WR; ws.merge_cells("A2:I2"); ws.row_dimensions[2].height=30
rows=[
(1,"Evidence","Source review + repository","Project memory, source docs 00–12",D(13),D(13),"Claude","Done",""),
(2,"Evidence","Independent reruns TC01–TC12 (3 runs each)","PRs #70–#81, 37 runs, artefacts",D(13),D(13),"Claude + AB","Done","AB downloaded artefacts"),
(3,"Evidence","Verified results matrix + decisions D2–D11","Results workbook: Verified Matrix, Decisions & Caveats",D(13),D(13),"Claude + AB","Done","Decisions agreed by AB"),
(4,"Evidence","TC13–TC15 added and verified","Ground truth (commit cbda52c), PRs #82–#84, 9 runs",D(13),D(13),"Claude + AB","Done",""),
(5,"Evidence","AI quality scoring (46 runs)","AI Quality Scoring, Scoring Rules, Inter-rater Check sheets",D(13),D(15),"Claude + AB","In progress","AB blind inter-rater scores pending"),
(6,"Evidence","Metrics, timing, remediation validation","Metrics, Timing, Remediation Validation sheets; PRs #85–#100",D(13),D(13),"Claude","Done",""),
(7,"Evidence","Evidence register + figures","Evidence Register (36 claims), 6 figures at 300 dpi",D(13),D(13),"Claude","Done",""),
(8,"Dissertation","Ch3 Practical Research Methodology (2,100 words)","Chapter 3 draft",D(13),D(14),"Claude","Next",""),
(9,"Dissertation","Ch4 Analysis of Results (2,500 words)","Chapter 4 draft with tables/figures",D(15),D(16),"Claude","To do","Needs inter-rater + screenshots"),
(10,"Dissertation","Ch2 Literature Review (2,700 words)","Chapter 2 draft (from A1 review)",D(17),D(18),"Claude","To do",""),
(11,"Dissertation","Ch5 Project Evaluation and Reflection (1,100 words)","Chapter 5 draft",D(19),D(19),"Claude + AB","To do","Needs AB reflection bullets, ethics status"),
(12,"Dissertation","Ch6 (550), Ch1 (800), Abstract (250)","Chapters 6, 1 and abstract",D(20),D(20),"Claude","To do",""),
(13,"Dissertation","References + Appendices A–G; full draft to AB","Complete draft in template",D(21),D(21),"Claude","To do","Needs template .docx, PROM04 proposal file"),
(14,"Dissertation","AB review + revisions","Revised draft",D(22),D(22),"AB + Claude","To do",""),
(15,"Presentation","Slides (max 10) + timed script","Slide deck + speaker script (12–13 min)",D(22),D(23),"Claude","To do",""),
(16,"Presentation","Record presentation (10–15 min)","MP4/AVI/PPTX recording",D(24),D(24),"AB","To do",""),
(17,"QA","Final QA: consistency, references, word count, format, PDF","QA checklist, final PDF",D(24),D(25),"Claude + AB","To do",""),
(18,"Submission","Submit dissertation PDF + recording","Turnitin submission",D(25),D(25),"AB","To do","Deadline 26 Oct 23:59 (confirm)"),
(19,"Close-out","Close test PRs #56–#68, #70–#100 unmerged","Clean repository",D(26),D(27),"Claude (AB approval)","To do","After submission"),
]
last=table(ws,4,["#","Phase","Milestone","Deliverable","Start","Finish","Owner","Status","Notes"],[5,13,44,44,16,16,14,12,36],rows,datecols=(5,6),statuscol=8,inputcols=(5,6,7,8,9))
s=last+2
ws.cell(row=s,column=1,value="Summary").font=TI
for k,(lab,f) in enumerate([("Milestones",f"=COUNTA(C5:C{last})"),("Done",f'=COUNTIF(H5:H{last},"Done")'),("In progress / next",f'=COUNTIF(H5:H{last},"In progress")+COUNTIF(H5:H{last},"Next")'),
    ("To do",f'=COUNTIF(H5:H{last},"To do")'),("Complete (%)",f"=C{s+2}/C{s+1}"),("Days to target submission (from today)",f'=DATE(2026,10,25)-TODAY()')],s+1):
    ws.cell(row=k,column=1,value=lab).font=BB; ws.merge_cells(start_row=k,start_column=1,end_row=k,end_column=2); c=ws.cell(row=k,column=3,value=f)
    if lab.startswith("Complete"): c.number_format="0%"
ws.freeze_panes="D5"
# ---- Deliverables
d=wb.create_sheet("Deliverables")
d["A1"]="Submission deliverables"; d["A1"].font=TI
rows=[("Dissertation","PDF via Turnitin; official template","9,000–11,000 words incl. abstract; excludes references, figures, tables, appendices",D(25),"To do"),
("Recorded presentation","MP4/AVI/PPTX","10–15 minutes; max 10 slides",D(25),"To do"),
("Results workbook","XLSX (supporting evidence, appendix E source)","Verified matrix, metrics, scoring, register",D(13),"Done"),
("Evidence archive","GitHub branch prom06-evidence","Raw artefacts, scans, PR comments, logs, figures, scripts",D(13),"Done"),
("Figures","PNG, 300 dpi","6 figures (architecture, detection matrix, rates, AI quality, timing, remediation)",D(13),"Done"),
("Screenshots","PNG from GitHub","8 screenshots (S1–S8)",D(16),"Pending AB")]
table(d,3,["Deliverable","Format","Requirement / content","Due","Status"],[24,36,60,16,12],rows,datecols=(4,),statuscol=5,inputcols=(4,5))
# ---- AB inputs
a=wb.create_sheet("AB Inputs")
a["A1"]="Inputs needed from AB"; a["A1"].font=TI
rows=[("TC13–TC15 artefacts","Verify metadata, confirm results, score runs",D(13),"Done"),
("terraform validate log (run 38412955313)","Evidence for invalid AI tag fix",D(13),"Done"),
("Official template (.docx)","Write chapters directly into the template",D(14),"Pending AB"),
("Confirm deadline on live Canvas","Schedule depends on it",D(14),"Pending AB"),
("Inter-rater scores (10 runs, blind)","Report scorer agreement in Ch4/Ch5",D(15),"Pending AB"),
("8 screenshots (S1–S8)","Figures in Ch4, Appendix F",D(16),"Pending AB"),
("Reflection bullets","Ch5.4 personal reflection — cannot be invented",D(18),"Pending AB"),
("Ethics approval status / reference","Ch1.6, Ch3.9, Ch5.5",D(19),"Pending AB"),
("Original PROM04 proposal file","Appendix A",D(19),"Pending AB"),
("Confluence fix yes/no","Corrects 'independent' wording and TC11 heading",None,"Pending AB")]
table(a,3,["Item","Why it is needed","Needed by","Status"],[40,50,16,12],rows,datecols=(3,),statuscol=4,inputcols=(3,4))
# ---- Word budget
w=wb.create_sheet("Word Budget")
w["A1"]="Word allocation (agreed 13 Oct 2026)"; w["A1"].font=TI
rows=[("Abstract",250,None),("Ch1 Introduction",800,None),("Ch2 Literature Review",2700,None),("Ch3 Practical Research Methodology",2100,None),
("Ch4 Analysis of Results",2500,None),("Ch5 Project Evaluation and Reflection",1100,None),("Ch6 Conclusion and Recommendations",550,None)]
table(w,3,["Section","Target words","Actual words","Difference"],[40,14,14,12],rows,inputcols=(2,3))
for r in range(4,11): w.cell(row=r,column=4,value=f'=IF(C{r}="","",C{r}-B{r})').border=BOR
w.cell(row=11,column=1,value="Total").font=BB; w["B11"]="=SUM(B4:B10)"; w["C11"]="=SUM(C4:C10)"; w["D11"]='=IF(C11=0,"",C11-B11)'
w.cell(row=12,column=1,value="Allowed range 9,000–11,000 (incl. abstract). Fill 'Actual words' as chapters are drafted.").font=IT
wb.save("/home/claude/prom06-outputs/PROM06_Project_Plan.xlsx")
