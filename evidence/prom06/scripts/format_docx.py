"""Apply PROM06 layout rules to the pandoc output: A4, 2.54 cm margins, 12 pt clear font, 1.5 spacing
(single in tables, code and contents), page numbers in the footer, bordered tables, centred title page."""
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

src, dst = sys.argv[1], sys.argv[2]
d = Document(src)
FONT = "Arial"

def set_font(style, size=None, bold=None, italic=None, color=None, name=FONT):
    f = style.font; f.name = name
    rpr = style.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rf.set(qn(a), name)
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if color is not None: f.color.rgb = color

styles = {s.name: s for s in d.styles}
for n in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text", "Bibliography"):
    if n in styles:
        set_font(styles[n], 12)
        pf = styles[n].paragraph_format
        pf.line_spacing = 1.0 if n == "Compact" else 1.5
        pf.space_after = Pt(6)
for n, sz in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12), ("Heading 4", 12), ("Title", 20)):
    if n in styles:
        set_font(styles[n], sz, bold=True, color=RGBColor(0, 0, 0))
        styles[n].paragraph_format.space_before = Pt(18 if n == "Heading 1" else 12)
        styles[n].paragraph_format.space_after = Pt(6)
        styles[n].paragraph_format.line_spacing = 1.15
        styles[n].paragraph_format.keep_with_next = True
for n in ("Image Caption", "Table Caption", "Caption"):
    if n in styles:
        set_font(styles[n], 10, italic=True, color=RGBColor(0, 0, 0))
        styles[n].paragraph_format.line_spacing = 1.0
        styles[n].paragraph_format.space_after = Pt(10)
for n in ("Source Code",):
    if n in styles:
        set_font(styles[n], 8, name="Courier New"); styles[n].paragraph_format.line_spacing = 1.0
if "Verbatim Char" in styles: set_font(styles["Verbatim Char"], 9, name="Courier New")
for n in [s for s in styles if s.startswith("toc") or s.startswith("TOC")]:
    styles[n].paragraph_format.line_spacing = 1.0

# title block
if "TitleBlock" in styles:
    st = styles["TitleBlock"]; set_font(st, 14); st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st.paragraph_format.space_after = Pt(14); st.paragraph_format.line_spacing = 1.15
first_title = True
for p in d.paragraphs:
    if p.style.name == "TitleBlock":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if first_title:
            for r in p.runs: r.font.size = Pt(22); r.font.bold = True
            first_title = False
    if p.style.name == "Title":  # pandoc metadata title duplicates the title page; remove
        p._element.getparent().remove(p._element)

# tables: borders, single spacing, 10 pt
def borders(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = OxmlElement(f'w:{e}'); x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), '4'); x.set(qn('w:color'), '999999'); b.append(x)
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None: tblPr.remove(old)
    tblPr.append(b)
for t in d.tables:
    borders(t)
    for i, row in enumerate(t.rows):
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.size = Pt(9 if len(t.columns) > 6 else 10)
                    if i == 0: r.font.bold = True

# page setup and footer page numbers
for s in d.sections:
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.54)
    fp = s.footer.paragraphs[0] if s.footer.paragraphs else s.footer.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for k, v in (("begin", None), (None, " PAGE "), ("end", None)):
        r = fp.add_run()
        if k:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), k); r._r.append(fc)
        else:
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = v; r._r.append(it)
        r.font.size = Pt(10); r.font.name = FONT
d.save(dst)
print("formatted")
